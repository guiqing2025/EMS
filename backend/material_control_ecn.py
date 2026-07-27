"""永联 ECR/ECN 变更单（.xls/.xlsx）解析，并按机型匹配在制订单。"""
from __future__ import annotations

import io
import re
from datetime import datetime
from typing import Any, Optional

from sqlalchemy.orm import Session

from engineering_service import normalize_code
from models import MaterialControl, SrmOrder

_MODEL_CODE_RE = re.compile(r"\b\d{2}\.\d{2,4}\.\d{4,}\b")
_CHECKED = ("■", "☑", "√", "✓", "▇", "✕", "x", "X", "true", "1", "是")


def _cell_str(value: Any) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, float) and value == int(value):
        return str(int(value))
    return str(value).strip()


def _is_checked(value: Any) -> bool:
    text = _cell_str(value)
    if not text:
        return False
    if text.lower() in ("true", "1", "yes"):
        return True
    return any(mark in text for mark in _CHECKED)


def _to_float(value: Any) -> float:
    text = _cell_str(value).replace(",", "")
    if not text:
        return 0.0
    try:
        return float(text)
    except ValueError:
        return 0.0


def _sheet_rows_xlrd(content: bytes) -> list[list[Any]]:
    import xlrd

    book = xlrd.open_workbook(file_contents=content)
    sheet = None
    for name in book.sheet_names():
        if "ecr" in name.lower() or "ecn" in name.lower():
            sheet = book.sheet_by_name(name)
            break
    if sheet is None:
        sheet = book.sheet_by_index(0)
    rows: list[list[Any]] = []
    for r in range(sheet.nrows):
        rows.append([sheet.cell_value(r, c) for c in range(sheet.ncols)])
    return rows


def _sheet_rows_xlsx(content: bytes) -> list[list[Any]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    ws = None
    for name in wb.sheetnames:
        if "ecr" in name.lower() or "ecn" in name.lower():
            ws = wb[name]
            break
    if ws is None:
        ws = wb.active
    rows: list[list[Any]] = []
    for row in ws.iter_rows(values_only=True):
        rows.append(list(row))
    return rows


def _load_rows(filename: str, content: bytes) -> list[list[Any]]:
    name = (filename or "").lower()
    if name.endswith((".xlsx", ".xlsm")):
        try:
            return _sheet_rows_xlsx(content)
        except Exception:
            pass
    # .xls 或 xlsx 失败回退
    try:
        return _sheet_rows_xlrd(content)
    except Exception as exc:  # noqa: BLE001
        if name.endswith((".xlsx", ".xlsm")):
            raise ValueError(f"无法解析 Excel：{exc}") from exc
        try:
            return _sheet_rows_xlsx(content)
        except Exception as exc2:  # noqa: BLE001
            raise ValueError(f"无法解析 ECN 文件：{exc2}") from exc2


def _find_model_in_row(row: list[Any]) -> tuple[str, str]:
    """从表头信息行提取机型编码与规格型号。"""
    model_code = ""
    model_name = ""
    joined = " | ".join(_cell_str(c) for c in row)
    # 优先「变更BOM编码」右侧
    for i, cell in enumerate(row):
        text = _cell_str(cell)
        if "变更BOM" in text or "BOM编码" in text:
            for j in range(i + 1, min(i + 6, len(row))):
                m = _MODEL_CODE_RE.search(_cell_str(row[j]))
                if m:
                    model_code = m.group(0)
                    break
            if not model_code:
                m = _MODEL_CODE_RE.search(text)
                if m:
                    model_code = m.group(0)
        if "变更项规格" in text or "规格型号" in text:
            for j in range(i + 1, min(i + 5, len(row))):
                cand = _cell_str(row[j])
                if cand and not cand.startswith("日期") and "变更" not in cand:
                    model_name = cand
                    break
    if not model_code:
        m = _MODEL_CODE_RE.search(joined)
        if m:
            model_code = m.group(0)
    return model_code, model_name


def _find_date(rows: list[list[Any]]) -> str:
    for row in rows[:8]:
        for i, cell in enumerate(row):
            if "日期" in _cell_str(cell):
                for j in range(i + 1, min(i + 4, len(row))):
                    raw = _cell_str(row[j]).replace("/", ".").replace("-", ".")
                    if re.match(r"\d{4}\.\d{1,2}\.\d{1,2}", raw):
                        parts = raw.split(".")
                        return f"{int(parts[0]):04d}-{int(parts[1]):02d}-{int(parts[2]):02d}"
            raw = _cell_str(cell).replace("/", ".").replace("-", ".")
            if re.fullmatch(r"\d{4}\.\d{1,2}\.\d{1,2}", raw):
                parts = raw.split(".")
                return f"{int(parts[0]):04d}-{int(parts[1]):02d}-{int(parts[2]):02d}"
    return datetime.utcnow().strftime("%Y-%m-%d")


def _extract_purchase_nos_from_text(text: str) -> list[str]:
    """从表内「订单号」区域尽量抽取采购订单号。"""
    found: list[str] = []
    # 永联：01-202607-CG-0320
    for m in re.finditer(r"\b\d{2}-\d{6}-[A-Za-z]{2}-\d{3,}\b", text):
        found.append(m.group(0))
    # 菲利斯式工单兜底
    for m in re.finditer(r"\b\d{4}-\d{8,}\b", text):
        found.append(m.group(0))
    # 去重保序
    out: list[str] = []
    seen: set[str] = set()
    for p in found:
        if p not in seen:
            seen.add(p)
            out.append(p)
    return out


def _find_detail_header(rows: list[list[Any]]) -> int:
    for i, row in enumerate(rows):
        texts = [_cell_str(c) for c in row]
        joined = "".join(texts)
        if "物料编码" in joined and ("增加" in joined or "删除" in joined or "变更" in joined):
            return i
    return -1


def _parse_detail_rows(rows: list[list[Any]], header_idx: int) -> list[dict]:
    changes: list[dict] = []
    for row in rows[header_idx + 1 :]:
        joined = "".join(_cell_str(c) for c in row)
        if not joined.strip():
            continue
        if "执行方式" in joined or "处理方式" in joined or "采购" in joined and "意见" in joined:
            break
        if "变更内容明细" in joined:
            continue

        is_add = _is_checked(row[0] if len(row) > 0 else "")
        is_change = _is_checked(row[1] if len(row) > 1 else "")
        is_delete = _is_checked(row[2] if len(row) > 2 else "")
        before_qty = _to_float(row[3] if len(row) > 3 else 0)
        after_qty = _to_float(row[4] if len(row) > 4 else 0)
        # 物料编码通常在 col5，也可能跨合并格
        code = ""
        for idx in (5, 6, 7):
            if idx < len(row):
                cand = _cell_str(row[idx])
                m = _MODEL_CODE_RE.search(cand) or re.search(r"\b\d{2}\.\d{2}\.\d+\b", cand)
                if m:
                    code = m.group(0)
                    break
                if cand and re.match(r"^[\d.]+$", cand) and len(cand) >= 8:
                    code = cand
                    break
        if not code:
            continue
        if not (is_add or is_change or is_delete):
            # 无勾选但有料号：按用量推断
            if after_qty > 0 and before_qty <= 0:
                is_add = True
            elif before_qty > 0 and after_qty <= 0:
                is_delete = True
            elif before_qty > 0 and after_qty > 0:
                is_change = True
            else:
                continue

        before_ref = _cell_str(row[12] if len(row) > 12 else "")
        after_ref = _cell_str(row[17] if len(row) > 17 else "")
        spec = _cell_str(row[8] if len(row) > 8 else "")
        desc = _cell_str(row[10] if len(row) > 10 else "")
        remark = " / ".join(x for x in (spec, desc) if x)[:200] or None

        if is_delete and not is_add and not is_change:
            changes.append(
                {
                    "control_qty": 0,
                    "remove_code": code,
                    "remove_qty": before_qty or 1,
                    "remove_refdes": before_ref or after_ref or None,
                    "add_code": "",
                    "add_qty": 0,
                    "add_refdes": None,
                    "remark": remark,
                    "action": "删除",
                }
            )
        elif is_add and not is_delete and not is_change:
            changes.append(
                {
                    "control_qty": 0,
                    "remove_code": "",
                    "remove_qty": 0,
                    "remove_refdes": None,
                    "add_code": code,
                    "add_qty": after_qty or 1,
                    "add_refdes": after_ref or before_ref or None,
                    "remark": remark,
                    "action": "增加",
                }
            )
        else:
            # 变更：同料号改用量/位号 → 删旧加新
            changes.append(
                {
                    "control_qty": 0,
                    "remove_code": code,
                    "remove_qty": before_qty if before_qty > 0 else 1,
                    "remove_refdes": before_ref or None,
                    "add_code": code,
                    "add_qty": after_qty if after_qty > 0 else 1,
                    "add_refdes": after_ref or before_ref or None,
                    "remark": remark,
                    "action": "变更",
                }
            )
    return changes


def _model_from_filename(filename: str) -> str:
    m = _MODEL_CODE_RE.search(filename or "")
    return m.group(0) if m else ""


def match_open_orders_for_model(db: Session, model_code: str) -> list[dict]:
    needle = normalize_code(model_code)
    if not needle:
        return []
    rows = (
        db.query(SrmOrder)
        .filter(SrmOrder.is_completed.is_(False))
        .order_by(SrmOrder.purchase_date.desc(), SrmOrder.id.desc())
        .all()
    )
    # 按采购订单号聚合
    by_pn: dict[str, dict] = {}
    for so in rows:
        if normalize_code(so.product_goods_no or "") != needle:
            continue
        pn = (so.purchase_no or "").strip()
        if not pn:
            continue
        qty = float(so.batch_pur_qty or so.output_qty or 0)
        hit = by_pn.get(pn)
        if not hit:
            by_pn[pn] = {
                "purchase_no": pn,
                "qty": qty,
                "customer_id": so.customer_id or "",
                "customer_name": so.customer_name or "",
                "product_goods_no": so.product_goods_no or "",
                "product_goods_name": so.product_goods_name or "",
                "selected": True,  # 默认全选
            }
        else:
            hit["qty"] = max(float(hit.get("qty") or 0), qty)
    return list(by_pn.values())


def parse_yonglian_ecn_file(
    db: Session,
    filename: str,
    content: bytes,
) -> dict:
    if not content:
        raise ValueError("空文件")
    rows = _load_rows(filename, content)
    if not rows:
        raise ValueError("表格无内容")

    title_ok = any("变更" in "".join(_cell_str(c) for c in row) and ("ECR" in "".join(_cell_str(c) for c in row) or "替代" in "".join(_cell_str(c) for c in row) or "ECN" in (filename or "").upper()) for row in rows[:5])
    if not title_ok:
        # 宽松：有变更BOM编码也认
        title_ok = any("变更BOM" in "".join(_cell_str(c) for c in row) for row in rows[:6])
    if not title_ok:
        raise ValueError("无法识别为永联 ECR/ECN 变更单，请确认上传正确表格")

    model_code = ""
    model_name = ""
    for row in rows[:8]:
        mc, mn = _find_model_in_row(row)
        if mc and not model_code:
            model_code = mc
        if mn and not model_name:
            model_name = mn
    if not model_code:
        model_code = _model_from_filename(filename)
    if not model_code:
        raise ValueError("未找到变更 BOM 编码（机型料号）")

    ecn_date = _find_date(rows)
    header_idx = _find_detail_header(rows)
    if header_idx < 0:
        raise ValueError("未找到变更内容明细表头")
    changes = _parse_detail_rows(rows, header_idx)
    if not changes:
        raise ValueError("未解析到有效换料行")

    # 表内订单号（常为空）
    head_text = "\n".join(" ".join(_cell_str(c) for c in row) for row in rows[:16])
    explicit_pos = _extract_purchase_nos_from_text(head_text)

    matched = match_open_orders_for_model(db, model_code)
    warnings: list[str] = []

    if explicit_pos:
        # 表内有订单号：优先勾选这些；仍展示机型匹配到的其它在制单供勾选
        explicit_set = set(explicit_pos)
        for m in matched:
            m["selected"] = m["purchase_no"] in explicit_set
        for pn in explicit_pos:
            if not any(m["purchase_no"] == pn for m in matched):
                so = (
                    db.query(SrmOrder)
                    .filter(SrmOrder.purchase_no == pn)
                    .order_by(SrmOrder.id.desc())
                    .first()
                )
                matched.append(
                    {
                        "purchase_no": pn,
                        "qty": float(so.batch_pur_qty or so.output_qty or 0) if so else 0,
                        "customer_id": (so.customer_id if so else "") or "",
                        "customer_name": (so.customer_name if so else "") or "",
                        "product_goods_no": (so.product_goods_no if so else "") or "",
                        "product_goods_name": (so.product_goods_name if so else "") or "",
                        "selected": True,
                        "from_sheet": True,
                    }
                )
                if not so:
                    warnings.append(f"表内订单号 {pn} 系统中不存在")
                elif so.is_completed:
                    warnings.append(f"表内订单号 {pn} 已结案")
    else:
        # 默认：该机型全部在制订单勾选
        if not matched:
            warnings.append(f"机型 {model_code} 暂无在制订单，请手工指定采购订单号后再导入")
        else:
            for m in matched:
                m["selected"] = True

    reason_bits = ["永联ECR"]
    if "临时变更" in head_text or "临时" in (filename or ""):
        reason_bits.append("临时变更")
    if "试产" in head_text:
        reason_bits.append("试产")
    if "物料调整" in head_text or "物料变更" in head_text:
        reason_bits.append("物料变更")
    if model_name:
        reason_bits.append(model_name)

    date_compact = ecn_date.replace("-", "")
    model_tail = model_code.split(".")[-1] if "." in model_code else model_code[-6:]
    suggested_no = f"YL-ECN-{model_tail}-{date_compact}"

    return {
        "source": "yonglian_ecn",
        "filename": filename or "",
        "model_code": model_code,
        "model_name": model_name,
        "ecn_date": ecn_date,
        "suggested_control_no": suggested_no,
        "control_type": "永联ECN",
        "reason": " · ".join(reason_bits),
        "changes": changes,
        "matched_orders": matched,
        "explicit_purchase_nos": explicit_pos,
        "warnings": warnings,
        "message": f"识别机型 {model_code}，换料 {len(changes)} 行，匹配在制订单 {len(matched)} 个",
    }


def _unique_control_no(db: Session, base: str) -> str:
    no = (base or "").strip() or f"YL-ECN-{datetime.utcnow().strftime('%Y%m%d%H%M%S')}"
    if not db.query(MaterialControl).filter(MaterialControl.control_no == no).first():
        return no
    for i in range(2, 50):
        cand = f"{no}-{i}"
        if not db.query(MaterialControl).filter(MaterialControl.control_no == cand).first():
            return cand
    return f"{no}-{datetime.utcnow().strftime('%H%M%S')}"


def build_control_payload_from_ecn_preview(
    db: Session,
    preview: dict,
    *,
    selected_purchase_nos: Optional[list[str]] = None,
    control_no: str = "",
) -> dict:
    """将预览结果转为 create_control 入参。"""
    model_code = str(preview.get("model_code") or "").strip()
    changes = preview.get("changes") or []
    if not model_code or not changes:
        raise ValueError("预览数据缺少机型或换料行")

    matched = preview.get("matched_orders") or []
    if selected_purchase_nos is None:
        selected = [m for m in matched if m.get("selected")]
    else:
        want = {(p or "").strip() for p in selected_purchase_nos if (p or "").strip()}
        selected = [m for m in matched if m.get("purchase_no") in want]
        # 允许只传订单号、不在 matched 里
        have = {m["purchase_no"] for m in selected}
        for pn in want - have:
            so = (
                db.query(SrmOrder)
                .filter(SrmOrder.purchase_no == pn)
                .order_by(SrmOrder.id.desc())
                .first()
            )
            selected.append(
                {
                    "purchase_no": pn,
                    "qty": float(so.batch_pur_qty or so.output_qty or 0) if so else 0,
                }
            )

    if not selected:
        raise ValueError("请至少选择一个在制订单")

    orders = []
    for m in selected:
        pn = str(m.get("purchase_no") or "").strip()
        qty = float(m.get("qty") or 0)
        if qty <= 0:
            so = (
                db.query(SrmOrder)
                .filter(SrmOrder.purchase_no == pn, SrmOrder.is_completed.is_(False))
                .order_by(SrmOrder.id.desc())
                .first()
            )
            if so:
                qty = float(so.batch_pur_qty or so.output_qty or 0)
        orders.append(
            {
                "purchase_no": pn,
                "order_qty": qty,
                "control_qty": qty,  # 整单临时变更
            }
        )

    clean_changes = []
    for c in changes:
        clean_changes.append(
            {
                "control_qty": 0,
                "remove_code": c.get("remove_code") or "",
                "remove_qty": float(c.get("remove_qty") or 0),
                "remove_refdes": c.get("remove_refdes"),
                "add_code": c.get("add_code") or "",
                "add_qty": float(c.get("add_qty") or 0),
                "add_refdes": c.get("add_refdes"),
                "remark": c.get("remark"),
            }
        )

    no = _unique_control_no(db, control_no or preview.get("suggested_control_no") or "")
    return {
        "control_no": no,
        "control_type": preview.get("control_type") or "永联ECN",
        "reason": preview.get("reason"),
        "ecn_no": no,
        "groups": [
            {
                "model_code": model_code,
                "orders": orders,
                "changes": clean_changes,
            }
        ],
    }
