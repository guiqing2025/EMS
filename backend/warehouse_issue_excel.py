"""按订单发料 Excel 模板导出与导入解析。"""
from __future__ import annotations

import io
import re
from typing import Any, Optional

from openpyxl import Workbook, load_workbook

from warehouse_inbound_excel import CODE_HEADERS, REMARK_HEADERS, _find_col, _float_val, _norm_header

# 不用泛化「数量」：齐套导出等表有「替代库存数量」会误匹配
ISSUE_QTY_HEADERS = ("本次发料", "发料数量", "发料数", "issue_qty", "qty", "Qty")
DEPT_HEADERS = ("部门", "department", "Dept")

ISSUE_TEMPLATE_HEADERS = [
    "物料编码",
    "品名",
    "需求数",
    "已发数",
    "待发数",
    "可用库存",
    "本次发料",
    "部门",
    "备注",
]


def _safe_filename_part(text: str) -> str:
    return re.sub(r'[\\/:*?"<>|]', "_", (text or "").strip()) or "order"


def build_issue_template_workbook(preview: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    order_no = preview.get("purchase_no") or ""
    model = preview.get("product_goods_no") or ""
    ws.title = "发料导入"[:31]
    ws.append([f"订单号：{order_no}", f"机型：{model}", f"订单量：{preview.get('order_qty') or 0}"])
    ws.append([])
    ws.append(ISSUE_TEMPLATE_HEADERS)
    for line in preview.get("lines") or []:
        dept = "dip" if (line.get("process") or "").lower() in ("dip", "assy") else "smt"
        suggested = float(line.get("suggested_qty") or 0)
        ws.append([
            line.get("material_code") or "",
            line.get("material_name") or "",
            line.get("required_qty") or 0,
            line.get("issued_qty") or 0,
            line.get("remain_qty") or 0,
            line.get("available_qty") or 0,
            suggested if suggested > 0 else "",
            dept,
            "",
        ])
    return wb


def issue_template_filename(preview: dict) -> str:
    order_no = _safe_filename_part(preview.get("purchase_no") or "order")
    model = _safe_filename_part(preview.get("product_goods_no") or "")
    return f"发料_{order_no}_{model}.xlsx"


def parse_issue_excel(content: bytes, preview: dict) -> dict:
    """解析发料 Excel，按 preview 中的料号匹配 material_id。"""
    try:
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    except Exception as exc:
        raise ValueError(f"无法读取 Excel 文件：{exc}") from exc

    code_map: dict[str, dict] = {}
    for line in preview.get("lines") or []:
        code = _norm_header(line.get("material_code"))
        if code:
            code_map[code.lower()] = line

    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        raise ValueError("Excel 为空")

    header_idx = -1
    code_col = qty_col = dept_col = remark_col = None

    for idx, row in enumerate(rows[:20]):
        cells = [_norm_header(c) for c in row]
        if not any(cells):
            continue
        cc = _find_col(cells, CODE_HEADERS)
        qc = _find_col(cells, ISSUE_QTY_HEADERS)
        # 标题行（订单号/机型）或空表头格一律跳过，避免误当表头
        if cc is None or qc is None:
            continue
        code_h = cells[cc]
        qty_h = cells[qc]
        if not code_h or not qty_h:
            continue
        if code_h.startswith("订单号") or "机型" in code_h:
            continue
        header_idx = idx
        code_col = cc
        qty_col = qc
        dept_col = _find_col(cells, DEPT_HEADERS)
        remark_col = _find_col(cells, REMARK_HEADERS)
        break

    if code_col is None or qty_col is None:
        raise ValueError("未找到表头，请使用系统下载的发料模板，或确保含「物料编码」「本次发料」列")

    items = []
    errors = []
    for row_no, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        if not row:
            continue
        cells = list(row)
        if code_col >= len(cells):
            continue
        code = _norm_header(cells[code_col])
        if not code or code in CODE_HEADERS or code.startswith("订单号"):
            continue
        qty = _float_val(cells[qty_col] if qty_col < len(cells) else None)
        if qty is None:
            continue
        dept_raw = _norm_header(cells[dept_col]) if dept_col is not None and dept_col < len(cells) else ""
        dept = "dip" if "dip" in dept_raw.lower() else "smt"
        remark = ""
        if remark_col is not None and remark_col < len(cells):
            remark = _norm_header(cells[remark_col])

        matched = code_map.get(code.lower())
        if not matched:
            for key, val in code_map.items():
                if key == code.lower() or val.get("material_code", "").lower() == code.lower():
                    matched = val
                    break
        if not matched:
            errors.append({"row": row_no, "material_code": code, "error": "料号不在本订单 BOM 清单中"})
            continue
        if not matched.get("material_id"):
            errors.append({"row": row_no, "material_code": code, "error": "物料未建档"})
            continue
        items.append({
            "material_id": matched["material_id"],
            "material_code": matched.get("material_code") or code,
            "qty": qty,
            "department": dept,
            "remark": remark,
        })

    if not items:
        if errors:
            sample = "；".join(f"第{e['row']}行 {e['material_code']}:{e['error']}" for e in errors[:3])
            raise ValueError(f"未解析到有效发料行（{len(errors)} 行无法匹配）。{sample}")
        raise ValueError("未解析到有效发料行，请在模板「本次发料」列填写大于 0 的数量后导入")

    return {
        "lines": items,
        "errors": errors,
        "parsed_count": len(items),
        "header_row": header_idx + 1,
    }
