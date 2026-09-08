# -*- coding: utf-8 -*-
"""工程资料自查：发料用量误用 + 全量 BOM 漏料对比。"""
from __future__ import annotations

import json
from collections import defaultdict
from datetime import datetime
from io import BytesIO
from typing import Any, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import Session

from eng_bom_import_service import analyze_import_bom_integrity
from engineering_service import normalize_code as norm_code
from models import BomLine, BomModel, EngIssuePrintLog, SrmOrder, SubstitutionRule

YONGLIAN_CUSTOMERS = {"yonglian", "a067"}


def _approx(a: float, b: float, tol: float = 0.02) -> bool:
    return abs(float(a) - float(b)) <= tol


def _order_qty_map(db: Session) -> dict[str, float]:
    out: dict[str, float] = {}
    for po, qty in (
        db.query(SrmOrder.purchase_no, SrmOrder.output_qty)
        .filter(SrmOrder.purchase_no.isnot(None), SrmOrder.purchase_no != "")
        .all()
    ):
        key = str(po or "").strip()
        if not key:
            continue
        val = float(qty or 0)
        if key not in out or val > out[key]:
            out[key] = val
    return out


def _rule_matches_bom(rule: SubstitutionRule, bom: BomModel) -> bool:
    mc = norm_code(bom.model_code)
    if norm_code(rule.parent_code) and norm_code(rule.parent_code) == mc:
        return True
    rpo = str(rule.purchase_no or "").strip()
    bpo = str(bom.purchase_no or "").strip()
    return bool(rpo and bpo and rpo == bpo)


def audit_substitution_qty_issues(db: Session) -> list[dict[str, Any]]:
    """非永联：替代规则 qty 与 BOM 单台用量不一致（修复前发料单会误用）。"""
    order_qty = _order_qty_map(db)
    boms = (
        db.query(BomModel)
        .filter(BomModel.is_active.is_(True), BomModel.line_count > 0)
        .all()
    )
    rules = (
        db.query(SubstitutionRule)
        .filter(
            SubstitutionRule.qty.isnot(None),
            SubstitutionRule.qty > 0,
            SubstitutionRule.confirm_status != "pending",
        )
        .all()
    )
    rules_by_customer: dict[str, list[SubstitutionRule]] = defaultdict(list)
    for r in rules:
        cid = (r.customer_id or "").strip().lower()
        if cid in YONGLIAN_CUSTOMERS:
            continue
        rules_by_customer[r.customer_id or ""].append(r)

    issues: list[dict[str, Any]] = []
    seen: set[tuple] = set()

    for bom in boms:
        if (bom.customer_id or "").strip().lower() in YONGLIAN_CUSTOMERS:
            continue
        if (bom.internal_code or "").strip() == "A067":
            continue
        cust_rules = rules_by_customer.get(bom.customer_id or "", [])
        if not cust_rules:
            continue
        lines = (
            db.query(BomLine)
            .filter(BomLine.bom_model_id == bom.id, BomLine.is_active.is_(True))
            .all()
        )
        for line in lines:
            mat_n = norm_code(line.material_code)
            if not mat_n:
                continue
            qpb = float(line.qty_per or 0)
            for rule in cust_rules:
                if norm_code(rule.comp_code) != mat_n:
                    continue
                if not _rule_matches_bom(rule, bom):
                    continue
                rq = float(rule.qty or 0)
                if _approx(rq, qpb):
                    continue
                key = (bom.id, mat_n, rule.id)
                if key in seen:
                    continue
                seen.add(key)
                po = str(bom.purchase_no or "").strip()
                oq = order_qty.get(po, 0) or 0
                expected_total = round(qpb * oq, 4) if oq > 0 else None
                pattern = "unknown"
                if oq > 0 and expected_total is not None and _approx(rq, expected_total):
                    pattern = "rule_qty_is_total_demand"
                elif oq > 0 and _approx(rq, oq) and qpb == 1:
                    pattern = "rule_qty_equals_order_qty"
                issues.append(
                    {
                        "bom_id": bom.id,
                        "customer_id": bom.customer_id,
                        "internal_code": bom.internal_code,
                        "model_code": bom.model_code,
                        "purchase_no": bom.purchase_no,
                        "order_qty": oq,
                        "material_code": line.material_code,
                        "position": line.position,
                        "bom_qty_per": qpb,
                        "rule_id": rule.id,
                        "sub_code": rule.sub_code,
                        "rule_qty": rq,
                        "rule_purchase_no": rule.purchase_no,
                        "pattern": pattern,
                        "wrong_print_qty_per": rq,
                        "correct_qty_per": qpb,
                        "wrong_need": round(rq * oq, 2) if oq > 0 else None,
                        "correct_need": round(qpb * oq, 2) if oq > 0 else None,
                    }
                )
    issues.sort(key=lambda x: (x.get("purchase_no") or "", x.get("material_code") or ""))
    return issues


def audit_all_bom_leaks(db: Session) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    boms = (
        db.query(BomModel)
        .filter(BomModel.is_active.is_(True), BomModel.line_count > 0)
        .order_by(BomModel.customer_id, BomModel.purchase_no, BomModel.model_code)
        .all()
    )
    for bom in boms:
        chk = analyze_import_bom_integrity(db, bom)
        status = "ok"
        if not chk.get("has_snapshot"):
            status = "no_snapshot"
        elif chk.get("missing_in_bom"):
            status = "leak"
        elif chk.get("extra_in_bom"):
            status = "extra_only"
        rows.append(
            {
                "bom_id": bom.id,
                "customer_id": bom.customer_id,
                "internal_code": bom.internal_code,
                "model_code": bom.model_code,
                "purchase_no": bom.purchase_no,
                "line_count": bom.line_count,
                "review_status": bom.eng_review_status,
                "source_file": bom.source_file,
                "status": status,
                "import_count": chk.get("import_count", 0),
                "active_count": chk.get("active_count", 0),
                "missing": list(chk.get("missing_in_bom") or []),
                "extra": list(chk.get("extra_in_bom") or []),
                "detail": chk.get("detail") or "",
            }
        )
    return rows


def audit_wrong_print_logs(db: Session) -> list[dict[str, Any]]:
    bad: list[dict[str, Any]] = []
    logs = db.query(EngIssuePrintLog).order_by(EngIssuePrintLog.id.desc()).all()
    for log in logs:
        if not log.materials_json:
            continue
        try:
            mats = json.loads(log.materials_json)
        except (json.JSONDecodeError, TypeError):
            continue
        if not isinstance(mats, list):
            continue
        oq = float(log.order_qty or 0) or 1
        for m in mats:
            if not isinstance(m, dict):
                continue
            code = m.get("material_code") or ""
            per = float(m.get("qty_per") or 0)
            need = float(m.get("issue_qty") or 0)
            bl = (
                db.query(BomLine)
                .filter(
                    BomLine.bom_model_id == log.bom_model_id,
                    BomLine.is_active.is_(True),
                )
                .all()
            )
            bom_line = next((l for l in bl if norm_code(l.material_code) == norm_code(code)), None)
            if not bom_line:
                continue
            bom_per = float(bom_line.qty_per or 0)
            if per <= bom_per * 1.001:
                continue
            if oq > 1 and _approx(per, bom_per * oq, tol=max(0.05, bom_per * oq * 0.001)):
                bad.append(
                    {
                        "log_id": log.id,
                        "purchase_no": log.purchase_no,
                        "model_code": log.model_code,
                        "material_code": code,
                        "printed_qty_per": per,
                        "bom_qty_per": bom_per,
                        "order_qty": oq,
                        "printed_need": need,
                        "expected_need": round(bom_per * oq, 2),
                        "created_at": log.created_at.isoformat() if log.created_at else "",
                    }
                )
    return bad


def build_reprint_recommendations(
    db: Session,
    sub_issues: list[dict[str, Any]],
    wrong_logs: list[dict[str, Any]],
) -> list[dict[str, Any]]:
    """对曾错误打印的订单，生成建议重打发料单明细（正确用量）。"""
    target_pos = {str(x.get("purchase_no") or "").strip() for x in wrong_logs if x.get("purchase_no")}
    if not target_pos:
        return []

    by_po: dict[str, list[dict[str, Any]]] = defaultdict(list)
    for row in sub_issues:
        po = str(row.get("purchase_no") or "").strip()
        if po in target_pos:
            by_po[po].append(row)

    recs: list[dict[str, Any]] = []
    for po in sorted(target_pos):
        rows = by_po.get(po, [])
        if not rows:
            for wl in wrong_logs:
                if str(wl.get("purchase_no") or "").strip() != po:
                    continue
                recs.append(
                    {
                        "purchase_no": po,
                        "model_code": wl.get("model_code"),
                        "material_code": wl.get("material_code"),
                        "sub_code": "",
                        "position": "",
                        "bom_qty_per": wl.get("bom_qty_per"),
                        "order_qty": wl.get("order_qty"),
                        "correct_need": wl.get("expected_need"),
                        "wrong_printed_qty_per": wl.get("printed_qty_per"),
                        "wrong_printed_need": wl.get("printed_need"),
                        "action": "请 Ctrl+F5 后重打发料单",
                    }
                )
            continue
        seen_mat: set[str] = set()
        for row in rows:
            mat = str(row.get("material_code") or "")
            sub = str(row.get("sub_code") or "")
            key = f"{mat}|{sub}"
            if key in seen_mat:
                continue
            seen_mat.add(key)
            recs.append(
                {
                    "purchase_no": po,
                    "model_code": row.get("model_code"),
                    "material_code": mat,
                    "sub_code": sub,
                    "position": row.get("position"),
                    "bom_qty_per": row.get("bom_qty_per"),
                    "order_qty": row.get("order_qty"),
                    "correct_need": row.get("correct_need"),
                    "wrong_printed_qty_per": row.get("wrong_print_qty_per"),
                    "wrong_printed_need": row.get("wrong_need"),
                    "action": "请 Ctrl+F5 后重打发料单；核对仓库实发",
                }
            )
    return recs


def summarize_affected_orders(sub_issues: list[dict[str, Any]]) -> list[dict[str, Any]]:
    """按订单汇总受影响料号数。"""
    bucket: dict[str, dict[str, Any]] = {}
    for row in sub_issues:
        po = str(row.get("purchase_no") or "").strip() or "(机型级/无订单号)"
        if po not in bucket:
            bucket[po] = {
                "purchase_no": po if po != "(机型级/无订单号)" else "",
                "customer_id": row.get("customer_id"),
                "internal_code": row.get("internal_code"),
                "model_code": row.get("model_code"),
                "order_qty": row.get("order_qty"),
                "issue_line_count": 0,
                "bom_ids": set(),
                "sample_materials": [],
            }
        b = bucket[po]
        b["issue_line_count"] += 1
        b["bom_ids"].add(row.get("bom_id"))
        if len(b["sample_materials"]) < 5:
            b["sample_materials"].append(row.get("material_code"))
    out = []
    for po in sorted(bucket.keys()):
        b = bucket[po]
        out.append(
            {
                "purchase_no": b["purchase_no"] or po,
                "customer_id": b["customer_id"],
                "internal_code": b["internal_code"],
                "model_code": b["model_code"],
                "order_qty": b["order_qty"],
                "issue_line_count": b["issue_line_count"],
                "bom_count": len(b["bom_ids"]),
                "sample_materials": "、".join(filter(None, b["sample_materials"])),
                "note": "修复前发料单会把替代规则qty当单台用量",
            }
        )
    return out


def run_eng_data_audit(db: Session) -> dict[str, Any]:
    sub_issues = audit_substitution_qty_issues(db)
    leak_rows = audit_all_bom_leaks(db)
    wrong_logs = audit_wrong_print_logs(db)
    affected_orders = summarize_affected_orders(sub_issues)
    reprint_recs = build_reprint_recommendations(db, sub_issues, wrong_logs)

    return {
        "generated_at": datetime.utcnow().isoformat(),
        "substitution_qty_issues": sub_issues,
        "affected_orders": affected_orders,
        "bom_leak_all": leak_rows,
        "wrong_print_logs": wrong_logs,
        "reprint_recommendations": reprint_recs,
        "summary": {
            "substitution_qty_issue_count": len(sub_issues),
            "substitution_affected_orders": len(affected_orders),
            "substitution_affected_boms": len({x["bom_id"] for x in sub_issues}),
            "bom_total": len(leak_rows),
            "bom_ok": sum(1 for r in leak_rows if r["status"] == "ok"),
            "bom_extra_only": sum(1 for r in leak_rows if r["status"] == "extra_only"),
            "bom_leak": sum(1 for r in leak_rows if r["status"] == "leak"),
            "bom_no_snapshot": sum(1 for r in leak_rows if r["status"] == "no_snapshot"),
            "wrong_print_log_rows": len(wrong_logs),
            "reprint_order_count": len({r["purchase_no"] for r in reprint_recs}),
        },
    }


def _write_sheet(ws, columns: list[tuple[str, str, int]], rows: list[dict[str, Any]]) -> None:
    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    for col_idx, (_, title, width) in enumerate(columns, 1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    for row_idx, row in enumerate(rows, 2):
        for col_idx, (field, _, _) in enumerate(columns, 1):
            val = row.get(field, "")
            if isinstance(val, float):
                val = val if val == val else ""
            ws.cell(row=row_idx, column=col_idx, value=val if val is not None else "")
    ws.freeze_panes = "A2"


def build_eng_audit_xlsx(audit: dict[str, Any]) -> bytes:
    wb = Workbook()
    ws_sum = wb.active
    ws_sum.title = "汇总"
    s = audit.get("summary") or {}
    ws_sum["A1"] = "工程资料自查报告"
    ws_sum["A1"].font = Font(bold=True, size=14)
    ws_sum["A2"] = f"生成时间：{audit.get('generated_at', '')}"
    rows = [
        ("发料用量风险行数", s.get("substitution_qty_issue_count")),
        ("涉及订单数", s.get("substitution_affected_orders")),
        ("涉及BOM套数", s.get("substitution_affected_boms")),
        ("BOM总数", s.get("bom_total")),
        ("BOM漏料一致", s.get("bom_ok")),
        ("BOM漏料", s.get("bom_leak")),
        ("无导入快照", s.get("bom_no_snapshot")),
        ("历史错误打印条数", s.get("wrong_print_log_rows")),
        ("建议重打订单数", s.get("reprint_order_count")),
    ]
    for i, (k, v) in enumerate(rows, 4):
        ws_sum.cell(row=i, column=1, value=k)
        ws_sum.cell(row=i, column=2, value=v)
    ws_sum.column_dimensions["A"].width = 24
    ws_sum.column_dimensions["B"].width = 16

    ws_orders = wb.create_sheet("受影响订单")
    _write_sheet(
        ws_orders,
        [
            ("purchase_no", "采购订单号", 20),
            ("customer_id", "客户ID", 12),
            ("internal_code", "厂内代码", 10),
            ("model_code", "机型", 18),
            ("order_qty", "订单量", 10),
            ("issue_line_count", "风险料号数", 12),
            ("bom_count", "BOM套数", 10),
            ("sample_materials", "示例料号", 36),
            ("note", "说明", 40),
        ],
        audit.get("affected_orders") or [],
    )

    ws_detail = wb.create_sheet("用量风险明细")
    _write_sheet(
        ws_detail,
        [
            ("purchase_no", "订单号", 20),
            ("model_code", "机型", 18),
            ("material_code", "BOM料号", 18),
            ("sub_code", "替代料", 18),
            ("position", "位号", 14),
            ("bom_qty_per", "BOM单台用量", 12),
            ("rule_qty", "规则qty", 12),
            ("order_qty", "订单量", 10),
            ("wrong_need", "误打需求", 14),
            ("correct_need", "正确需求", 14),
            ("pattern", "规律", 24),
        ],
        audit.get("substitution_qty_issues") or [],
    )

    ws_reprint = wb.create_sheet("建议重打发料单")
    _write_sheet(
        ws_reprint,
        [
            ("purchase_no", "订单号", 20),
            ("model_code", "机型", 18),
            ("material_code", "BOM料号", 18),
            ("sub_code", "投产替代料", 18),
            ("position", "位号", 14),
            ("bom_qty_per", "正确单台用量", 14),
            ("order_qty", "订单量", 10),
            ("correct_need", "正确需求数", 14),
            ("wrong_printed_qty_per", "曾误打用量", 14),
            ("wrong_printed_need", "曾误打需求", 14),
            ("action", "建议操作", 28),
        ],
        audit.get("reprint_recommendations") or [],
    )

    leaks = [r for r in (audit.get("bom_leak_all") or []) if r.get("status") == "leak"]
    if leaks:
        ws_leak = wb.create_sheet("BOM漏料")
        _write_sheet(
            ws_leak,
            [
                ("customer_id", "客户", 12),
                ("purchase_no", "订单号", 20),
                ("model_code", "机型", 18),
                ("import_count", "导入数", 10),
                ("active_count", "生效数", 10),
                ("detail", "说明", 40),
            ],
            leaks,
        )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
