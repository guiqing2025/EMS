"""从共享盘「员工花名册」同步员工档案。"""

from __future__ import annotations

import logging
import re
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from config import get_hr_roster_path
from models import HrEmployee

logger = logging.getLogger(__name__)

SHEET_NAME_CANDIDATES = ("员工花名册 ", "员工花名册")
HEADER_ROW = 21
# G=序号 … V=备注 → 0-based index 6..21
COL = {
    "seq": 6,
    "employee_no": 7,
    "name": 8,
    "gender": 9,
    "id_card": 10,
    "age": 11,
    "age_band": 12,
    "phone": 13,
    "address": 14,
    "hire_date": 15,
    "tenure_years": 16,
    "tenure_band": 17,
    "department": 18,
    "position": 19,
    "education": 20,
    "remark": 21,
}


def get_roster_path() -> Path:
    return Path(get_hr_roster_path()).expanduser()


def roster_accessible(path: Optional[Path] = None) -> bool:
    p = path or get_roster_path()
    try:
        return p.is_file()
    except OSError:
        return False


def _cell(row: tuple, idx: int) -> Any:
    if idx < 0 or idx >= len(row):
        return None
    return row[idx]


def _text(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, float) and value == int(value):
        s = str(int(value))
    else:
        s = str(value).strip()
    if not s or s.upper() in {"#VALUE!", "#REF!", "#N/A", "NONE", "NULL"}:
        return None
    return s


def _number(value: Any) -> Optional[float]:
    text = _text(value)
    if text is None:
        return None
    try:
        return float(text)
    except ValueError:
        return None


def _hire_date(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = _text(value)
    if not text:
        return None
    # Excel serial as float already handled via datetime from data_only
    m = re.match(r"^(\d{4})[-/](\d{1,2})[-/](\d{1,2})", text)
    if m:
        y, mo, d = int(m.group(1)), int(m.group(2)), int(m.group(3))
        try:
            return date(y, mo, d).isoformat()
        except ValueError:
            return text[:10]
    return text[:16]


def _pick_sheet(wb):
    names = {n: n for n in wb.sheetnames}
    for cand in SHEET_NAME_CANDIDATES:
        if cand in names:
            return wb[cand]
    for n in wb.sheetnames:
        if "花名册" in n.replace(" ", ""):
            return wb[n]
    raise FileNotFoundError(f"未找到「员工花名册」工作表，现有: {wb.sheetnames}")


def parse_roster_rows(path: Path) -> list[dict]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = _pick_sheet(wb)
        rows: list[dict] = []
        seen: set[str] = set()
        for i, row in enumerate(ws.iter_rows(min_row=HEADER_ROW + 1, values_only=True), start=HEADER_ROW + 1):
            emp_no = _text(_cell(row, COL["employee_no"]))
            name = _text(_cell(row, COL["name"]))
            if not emp_no or not name:
                continue
            emp_no = emp_no.upper()
            if emp_no in seen:
                logger.warning("花名册重复工号跳过: %s @row %s", emp_no, i)
                continue
            seen.add(emp_no)
            rows.append(
                {
                    "employee_no": emp_no,
                    "name": name,
                    "gender": _text(_cell(row, COL["gender"])),
                    "id_card": _text(_cell(row, COL["id_card"])),
                    "age": _number(_cell(row, COL["age"])),
                    "age_band": _text(_cell(row, COL["age_band"])),
                    "phone": _text(_cell(row, COL["phone"])),
                    "address": _text(_cell(row, COL["address"])),
                    "hire_date": _hire_date(_cell(row, COL["hire_date"])),
                    "tenure_years": _number(_cell(row, COL["tenure_years"])),
                    "tenure_band": _text(_cell(row, COL["tenure_band"])),
                    "department": _text(_cell(row, COL["department"])),
                    "position": _text(_cell(row, COL["position"])),
                    "education": _text(_cell(row, COL["education"])),
                    "remark": _text(_cell(row, COL["remark"])),
                }
            )
        return rows
    finally:
        wb.close()


def _phone(value: Any) -> Optional[str]:
    text = _text(value)
    if not text:
        return None
    digits = re.sub(r"\D", "", text)
    return digits or None


def parse_dorm_rows(path: Path) -> list[dict]:
    """解析「宿舍人员明细」：当前住宿（无搬离时间）名单。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        if "宿舍人员明细" not in wb.sheetnames:
            logger.warning("花名册缺少「宿舍人员明细」工作表，跳过住宿同步")
            return []
        ws = wb["宿舍人员明细"]
        rows: list[dict] = []
        for i, row in enumerate(ws.iter_rows(min_row=5, max_col=8, values_only=True), start=5):
            name = _text(_cell(row, 2))
            room = _text(_cell(row, 1))
            if not name or not room:
                continue
            move_out = _hire_date(_cell(row, 7))
            if move_out:
                continue  # 已搬离，不算当前住宿
            rows.append(
                {
                    "name": name,
                    "dorm_room": room,
                    "phone": _phone(_cell(row, 4)),
                    "gender": _text(_cell(row, 3)),
                    "move_in": _hire_date(_cell(row, 6)),
                    "row": i,
                }
            )
        return rows
    finally:
        wb.close()


def _match_dorm_employee(
    item: dict,
    by_phone: dict[str, HrEmployee],
    by_name: dict[str, list[HrEmployee]],
) -> Optional[HrEmployee]:
    phone = item.get("phone")
    if phone and phone in by_phone:
        return by_phone[phone]
    name = item.get("name") or ""
    cands = by_name.get(name) or []
    if len(cands) == 1:
        return cands[0]
    if len(cands) > 1 and phone:
        for c in cands:
            if _phone(c.phone) == phone:
                return c
    return None


def apply_dorm_sync(db: Session, path: Path, *, now: Optional[datetime] = None) -> dict:
    """将宿舍人员明细同步到员工档案的 lives_in_dorm / dorm_room。"""
    now = now or datetime.utcnow()
    dorm_rows = parse_dorm_rows(path)
    employees = db.query(HrEmployee).all()
    by_phone: dict[str, HrEmployee] = {}
    by_name: dict[str, list[HrEmployee]] = {}
    for emp in employees:
        p = _phone(emp.phone)
        if p:
            by_phone[p] = emp
        key = (emp.name or "").strip()
        if key:
            by_name.setdefault(key, []).append(emp)

    matched_ids: set[int] = set()
    matched = 0
    unmatched: list[str] = []
    for item in dorm_rows:
        emp = _match_dorm_employee(item, by_phone, by_name)
        if not emp:
            unmatched.append(f"{item['name']}({item['dorm_room']})")
            continue
        emp.lives_in_dorm = True
        emp.dorm_room = item["dorm_room"]
        emp.updated_at = now
        matched_ids.add(emp.id)
        matched += 1

    cleared = 0
    for emp in employees:
        if emp.id in matched_ids:
            continue
        if emp.lives_in_dorm or emp.dorm_room:
            emp.lives_in_dorm = False
            emp.dorm_room = None
            emp.updated_at = now
            cleared += 1

    if unmatched:
        logger.warning("宿舍明细未匹配到员工档案: %s", "、".join(unmatched))

    return {
        "dorm_total": len(dorm_rows),
        "dorm_matched": matched,
        "dorm_cleared": cleared,
        "dorm_unmatched": unmatched,
    }


def sync_hr_roster(db: Session, path: Optional[Path] = None) -> dict:
    roster = path or get_roster_path()
    if not roster_accessible(roster):
        raise FileNotFoundError(
            f"共享盘花名册不可访问: {roster}。请确认已挂载「共享-测试软件资料」。"
        )

    parsed = parse_roster_rows(roster)
    if not parsed:
        raise RuntimeError(f"花名册未解析到员工行: {roster}")

    now = datetime.utcnow()
    existing = {e.employee_no: e for e in db.query(HrEmployee).all()}
    seen_nos: set[str] = set()
    created = 0
    updated = 0

    for data in parsed:
        no = data["employee_no"]
        seen_nos.add(no)
        row = existing.get(no)
        if row:
            for k, v in data.items():
                setattr(row, k, v)
            # 手工锁定在职状态时，同步不强制改回在职/离职
            if not bool(getattr(row, "status_locked", False)):
                row.is_active = True
                row.leave_date = None
            row.synced_at = now
            row.updated_at = now
            updated += 1
        else:
            db.add(
                HrEmployee(
                    **data,
                    is_active=True,
                    source="roster",
                    synced_at=now,
                    created_at=now,
                    updated_at=now,
                )
            )
            created += 1

    inactivated = 0
    for no, row in existing.items():
        if no not in seen_nos and row.is_active:
            if bool(getattr(row, "status_locked", False)):
                continue
            row.is_active = False
            if not row.leave_date:
                row.leave_date = now.date().isoformat()
            if not row.leave_reason:
                row.leave_reason = "花名册同步：已不在共享盘花名册"
            row.synced_at = now
            row.updated_at = now
            inactivated += 1

    db.flush()
    dorm_stat = apply_dorm_sync(db, roster, now=now)

    db.commit()
    active = db.query(HrEmployee).filter(HrEmployee.is_active.is_(True)).count()
    dorm_count = db.query(HrEmployee).filter(HrEmployee.lives_in_dorm.is_(True)).count()
    logger.info(
        "人事花名册同步完成: 新增 %s 更新 %s 离职标记 %s 在职 %s；住宿匹配 %s/%s 清除 %s",
        created,
        updated,
        inactivated,
        active,
        dorm_stat["dorm_matched"],
        dorm_stat["dorm_total"],
        dorm_stat["dorm_cleared"],
    )
    return {
        "path": str(roster),
        "parsed": len(parsed),
        "created": created,
        "updated": updated,
        "inactivated": inactivated,
        "active_count": active,
        "dorm_count": dorm_count,
        "dorm_matched": dorm_stat["dorm_matched"],
        "dorm_total": dorm_stat["dorm_total"],
        "dorm_cleared": dorm_stat["dorm_cleared"],
        "dorm_unmatched": dorm_stat["dorm_unmatched"],
        "synced_at": now.isoformat() + "Z",
    }


def list_departments(db: Session) -> list[str]:
    rows = (
        db.query(HrEmployee.department)
        .filter(HrEmployee.department.isnot(None), HrEmployee.department != "")
        .distinct()
        .order_by(HrEmployee.department)
        .all()
    )
    return [r[0] for r in rows if r[0]]


def mask_id_card(id_card: Optional[str]) -> Optional[str]:
    if not id_card:
        return id_card
    s = id_card.strip()
    if len(s) < 8:
        return s
    return s[:4] + "*" * (len(s) - 8) + s[-4:]
