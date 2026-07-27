from datetime import datetime
from io import BytesIO
from typing import Optional
from urllib.parse import quote

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from models import SrmOrder

EXPORT_COLUMNS = [
    ("customer_name", "客户", 10),
    ("purchase_no", "采购订单号", 18),
    ("purchase_seq", "行序号", 8),
    ("purchase_phase_seq", "行相序", 8),
    ("product_goods_no", "PCBA料号", 18),
    ("product_goods_name", "品名", 22),
    ("product_spec", "规格", 40),
    ("batch_pur_qty", "订单量", 10),
    ("delivery_qty", "已交货", 10),
    ("receive_qty", "已收货", 10),
    ("un_receive_qty", "未收货", 10),
    ("purchase_date", "下单日期", 12),
    ("expect_arrival_date", "交期", 12),
    ("srm_status_name", "SRM状态", 12),
    ("is_completed", "完成", 8),
    ("remark", "备注", 16),
    ("workorder_no", "SRM工单号", 18),
    ("order_type_name", "订单类型", 12),
    ("data_source", "数据来源", 14),
]


def build_orders_xlsx(orders: list[SrmOrder]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "订单列表"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    for col_idx, (_, title, width) in enumerate(EXPORT_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, order in enumerate(orders, 2):
        for col_idx, (field, _, _) in enumerate(EXPORT_COLUMNS, 1):
            value = getattr(order, field, "")
            if field == "is_completed":
                value = "是" if value else "否"
            elif isinstance(value, float):
                value = value if value else 0
            ws.cell(row=row_idx, column=col_idx, value=value if value is not None else "")

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_COLUMNS))}{max(len(orders) + 1, 1)}"

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


BOM_LINE_COLUMNS = [
    ("seq", "序号", 8),
    ("material_code", "元件料号", 18),
    ("material_name", "品名", 24),
    ("mount_type", "贴装", 8),
    ("mount_side", "面别", 8),
    ("spec", "规格", 32),
    ("qty_per", "用量", 10),
    ("unit", "单位", 8),
    ("position", "位号", 24),
    ("process", "工艺", 10),
    ("remark", "备注", 16),
]


def bom_export_filename(model_code: str) -> str:
    safe_code = (model_code or "bom").replace("/", "-").replace("\\", "-")
    return f"BOM_{safe_code}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"


def bom_export_content_disposition(model_code: str) -> str:
    filename = bom_export_filename(model_code)
    ascii_name = f"bom_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"


def _write_bom_sheet(ws, bom: dict, lines: list[dict], *, title: str, profile_info: Optional[dict] = None) -> None:
    ws.title = title[:31]

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    summary_rows = [
        ("内部代码", bom.get("internal_code") or ""),
        ("客户", bom.get("customer_name") or bom.get("customer_id") or ""),
        ("机型料号", bom.get("model_code") or ""),
        ("品名", bom.get("model_name") or ""),
        ("规格", bom.get("model_spec") or ""),
        ("行数", len(lines)),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M:%S")),
    ]
    if profile_info:
        summary_rows.insert(5, ("贴装画像", profile_info.get("profile") or ""))
        if profile_info.get("source") == "manual":
            summary_rows.insert(6, ("画像来源", "手工指定"))
        elif profile_info.get("confidence"):
            summary_rows.insert(6, ("画像置信", profile_info.get("confidence")))

    for row_idx, (label, value) in enumerate(summary_rows, 1):
        ws.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row_idx, column=2, value=value)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 40

    start_row = len(summary_rows) + 2
    for col_idx, (_, col_title, width) in enumerate(BOM_LINE_COLUMNS, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=col_title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    for row_idx, line in enumerate(lines, start_row + 1):
        for col_idx, (field, _, _) in enumerate(BOM_LINE_COLUMNS, 1):
            value = line.get(field, "")
            if isinstance(value, float):
                value = value if value else 0
            ws.cell(row=row_idx, column=col_idx, value=value if value is not None else "")

    if lines:
        ws.freeze_panes = f"A{start_row + 1}"
        ws.auto_filter.ref = (
            f"A{start_row}:{get_column_letter(len(BOM_LINE_COLUMNS))}{start_row + len(lines)}"
        )


def build_bom_xlsx(bom: dict, lines: list[dict], profile_info: Optional[dict] = None) -> bytes:
    wb = Workbook()
    _write_bom_sheet(wb.active, bom, lines, title="BOM明细", profile_info=profile_info)

    smt_lines = [line for line in lines if line.get("mount_type") == "SMT"]
    dip_lines = [line for line in lines if line.get("mount_type") == "DIP"]
    assy_lines = [line for line in lines if line.get("mount_type") == "ASSY"]
    other_lines = [
        line for line in lines if line.get("mount_type") not in ("SMT", "DIP", "ASSY")
    ]

    if smt_lines:
        ws_smt = wb.create_sheet("SMT")
        _write_bom_sheet(ws_smt, bom, smt_lines, title="SMT")
    if dip_lines:
        ws_dip = wb.create_sheet("DIP")
        _write_bom_sheet(ws_dip, bom, dip_lines, title="DIP")
    if assy_lines:
        ws_assy = wb.create_sheet("装配")
        _write_bom_sheet(ws_assy, bom, assy_lines, title="装配")
    if other_lines and (smt_lines or dip_lines or assy_lines):
        ws_other = wb.create_sheet("其他")
        _write_bom_sheet(ws_other, bom, other_lines, title="其他")

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()


def export_filename() -> str:
    return f"菲利斯订单_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"


def export_content_disposition() -> str:
    filename = export_filename()
    ascii_name = f"orders_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"


KITTING_STATUS_LABELS = {
    "ready": "齐套",
    "partial": "部分齐",
    "shortage": "欠料",
    "unknown": "未算",
    "unbound": "未绑BOM",
}

KITTING_LINE_COLUMNS = [
    ("material_code", "料号", 16),
    ("material_name", "品名", 24),
    ("mount_type", "贴装", 8),
    ("mount_side", "面别", 8),
    ("spec", "规格", 28),
    ("qty_per", "用量", 10),
    ("unit", "单位", 8),
    ("position", "位号", 20),
    ("required_qty", "需求", 10),
    ("excel_in_qty", "来料数", 10),
    ("stock_qty", "库存", 10),
    ("main_gap_qty", "本单欠", 10),
    ("ledger_owe_qty", "账本欠", 10),
    ("kitting_supply_qty", "齐套可用", 10),
    ("shortage_qty", "欠料", 10),
    ("status_label", "状态", 10),
    ("substitute_codes_text", "替代料号", 28),
    ("substitute_stock_qty_text", "替代库存数量", 16),
]


def kitting_export_filename(model_code: str, order_qty: float) -> str:
    safe_code = (model_code or "kitting").replace("/", "-").replace("\\", "-")
    return f"备料齐套_{safe_code}_{int(order_qty) if order_qty == int(order_qty) else order_qty}_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"


def kitting_export_content_disposition(model_code: str, order_qty: float) -> str:
    filename = kitting_export_filename(model_code, order_qty)
    ascii_name = f"kitting_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return f"attachment; filename=\"{ascii_name}\"; filename*=UTF-8''{quote(filename)}"


def build_kitting_xlsx(kit: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "备料齐套"

    header_font = Font(bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
    header_align = Alignment(horizontal="center", vertical="center")

    summary_rows = [
        ("机型料号", kit.get("model_code") or kit.get("product_goods_no") or ""),
        ("机型名称", kit.get("model_name") or ""),
        ("客户", kit.get("customer_name") or kit.get("customer_id") or ""),
        ("订单号", kit.get("purchase_no") or ""),
        ("订单量", kit.get("order_qty") or 0),
        ("整体状态", KITTING_STATUS_LABELS.get(kit.get("material_status"), kit.get("material_status") or "")),
        ("齐套行数", kit.get("ready_count", 0)),
        ("部分齐行数", kit.get("partial_count", 0)),
        ("欠料行数", kit.get("shortage_count", 0)),
        ("合计行数", kit.get("total_lines", 0)),
    ]
    for row_idx, (label, value) in enumerate(summary_rows, 1):
        ws.cell(row=row_idx, column=1, value=label).font = Font(bold=True)
        ws.cell(row=row_idx, column=2, value=value)
    ws.column_dimensions["A"].width = 12
    ws.column_dimensions["B"].width = 36

    start_row = len(summary_rows) + 2
    for col_idx, (_, title, width) in enumerate(KITTING_LINE_COLUMNS, 1):
        cell = ws.cell(row=start_row, column=col_idx, value=title)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = header_align
        ws.column_dimensions[get_column_letter(col_idx)].width = width

    lines = kit.get("lines") or []
    for row_idx, line in enumerate(lines, start_row + 1):
        subs = line.get("substitutes") or []
        row_data = {
            **line,
            "status_label": KITTING_STATUS_LABELS.get(line.get("status"), line.get("status") or ""),
            "substitute_codes_text": "、".join(line.get("substitute_codes") or []),
            "substitute_stock_qty_text": "；".join(
                str(int(s.get("stock_qty", s.get("available_qty", 0)))) if float(s.get("stock_qty", s.get("available_qty", 0)) or 0) % 1 == 0
                else str(s.get("stock_qty", s.get("available_qty", 0)))
                for s in subs
            ) if subs else "",
        }
        for col_idx, (field, _, _) in enumerate(KITTING_LINE_COLUMNS, 1):
            value = row_data.get(field, "")
            if isinstance(value, float):
                value = value if value else 0
            ws.cell(row=row_idx, column=col_idx, value=value if value is not None else "")

    if lines:
        ws.freeze_panes = f"A{start_row + 1}"
        ws.auto_filter.ref = (
            f"A{start_row}:{get_column_letter(len(KITTING_LINE_COLUMNS))}{start_row + len(lines)}"
        )

    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
