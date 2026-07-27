"""亿兰科（A120）BOM：SMT / DIP 两份 Excel 分导后合并为系统 BOM。"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Optional

from openpyxl import load_workbook

from bom_excel import (
    ParsedBom,
    ParsedBomLine,
    _cell_code,
    _cell_float,
    _cell_str,
    _header_map,
    _norm_header,
    _pick_col,
)

# 整机料号 3001-xxxx；贴片半成品 3003-xxxx
_RE_FINISHED = re.compile(r"(3001-\d+[A-Za-z]?)", re.I)
_RE_SMT_SEMI = re.compile(r"(3003-\d+[A-Za-z]?)", re.I)


# 用量列别名：生产用 sheet 为「用量」；ERP 导出 Sheet1 常为「用量:分子」
_YILANKE_QTY_HEADERS = ("用量", "用量:分子", "用量分子", "需求")


def is_yilanke_bom_headers(headers: dict[str, int]) -> bool:
    """亿兰科表头：子项物料编码 + 用量（区别于永联「子物料编码+数量」）。"""
    has_code = _pick_col(headers, "子项物料编码") is not None
    has_qty = _pick_col(headers, *_YILANKE_QTY_HEADERS) is not None
    return bool(has_code and has_qty)


def find_yilanke_header_row(rows: list, max_scan: int = 40) -> Optional[int]:
    for idx, row in enumerate(rows[:max_scan]):
        if is_yilanke_bom_headers(_header_map(row)):
            return idx
    return None


def _title_blob(rows: list, header_idx: int) -> str:
    parts: list[str] = []
    for row in rows[: max(header_idx, 0) + 1]:
        if not row:
            continue
        for cell in row:
            if cell is None:
                continue
            t = str(cell).strip()
            if t:
                parts.append(t)
    return " ".join(parts)


def extract_yilanke_meta(rows: list, header_idx: int, filename: str = "") -> dict:
    """从标题/文件名提取整机料号、半成品料号、品名线索。"""
    blob = _title_blob(rows, header_idx) + " " + (filename or "")
    finished = None
    smt_semi = None
    m_fin = _RE_FINISHED.search(blob)
    if m_fin:
        finished = m_fin.group(1).upper()
    m_smt = _RE_SMT_SEMI.search(blob)
    if m_smt:
        smt_semi = m_smt.group(1).upper()
    # 括号内整机：3003-0001B（3001-0028C）
    m_paren = re.search(r"3003-\d+[A-Za-z]?\s*[（(]\s*(3001-\d+[A-Za-z]?)", blob, re.I)
    if m_paren:
        finished = m_paren.group(1).upper()
        if not smt_semi:
            m0 = _RE_SMT_SEMI.search(blob)
            if m0:
                smt_semi = m0.group(1).upper()
    return {
        "finished_code": finished,
        "smt_semi_code": smt_semi,
        "title": blob.strip()[:200],
    }


def detect_yilanke_role(rows: list, header_idx: int, filename: str = "", lines: Optional[list] = None) -> str:
    """返回 smt | dip。"""
    meta = extract_yilanke_meta(rows, header_idx, filename)
    name = (Path(filename).stem if filename else "").upper()
    if _looks_like_dip_parent(lines or [], meta):
        return "dip"
    if name.startswith("3001") or (meta.get("finished_code") and not meta.get("smt_semi_code")):
        return "dip"
    if name.startswith("3003") or meta.get("smt_semi_code"):
        return "smt"
    if lines:
        smt_kw = sum(
            1
            for ln in lines
            if any(k in ((ln.material_name or "") + (ln.spec or "")) for k in ("贴片", "SMD", "PCB空板"))
        )
        if smt_kw >= max(3, len(lines) // 3):
            return "smt"
    return "dip"


def _looks_like_dip_parent(lines: list, meta: dict) -> bool:
    semi = (meta.get("smt_semi_code") or "").upper()
    for ln in lines or []:
        code = (ln.material_code or "").upper()
        name = ln.material_name or ""
        if semi and code == semi:
            return True
        if "半成品" in name or "SMD半成品" in name.upper().replace(" ", ""):
            return True
        if code.startswith("3003-"):
            return True
    return False


def parse_yilanke_rows(
    rows: list,
    header_idx: int,
    *,
    role: str = "",
    filename: str = "",
) -> ParsedBom:
    headers = _header_map(rows[header_idx])
    if not is_yilanke_bom_headers(headers):
        raise ValueError("非亿兰科 BOM 格式（需含：子项物料编码 + 用量）")

    idx_seq = _pick_col(headers, "项次")
    idx_code = _pick_col(headers, "子项物料编码")
    idx_name = _pick_col(headers, "子项物料名称")
    idx_unit = _pick_col(headers, "单位", "子项单位")
    idx_qty = _pick_col(headers, *_YILANKE_QTY_HEADERS)
    idx_pos = _pick_col(headers, "位置号", "位号")
    idx_remark = _pick_col(headers, "备注", "物料备注")

    meta = extract_yilanke_meta(rows, header_idx, filename)
    lines: list[ParsedBomLine] = []
    sort_i = 0
    for row in rows[header_idx + 1 :]:
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        code = _cell_code(row[idx_code]) if idx_code is not None and idx_code < len(row) else None
        if not code:
            continue
        # 跳过表头重复
        if _norm_header(code) in ("子项物料编码", "项次"):
            continue
        name = _cell_str(row[idx_name]) if idx_name is not None and idx_name < len(row) else None
        unit = _cell_str(row[idx_unit]) if idx_unit is not None and idx_unit < len(row) else None
        qty = _cell_float(row[idx_qty], 1.0) if idx_qty is not None and idx_qty < len(row) else 1.0
        pos = _cell_str(row[idx_pos]) if idx_pos is not None and idx_pos < len(row) else None
        remark = _cell_str(row[idx_remark]) if idx_remark is not None and idx_remark < len(row) else None
        seq = _cell_str(row[idx_seq]) if idx_seq is not None and idx_seq < len(row) else None
        sort_i += 1
        lines.append(
            ParsedBomLine(
                seq=seq,
                material_code=code,
                material_name=name,
                spec=None,
                unit=(unit or "PCS").upper() if unit else "PCS",
                qty_per=qty if qty > 0 else 1.0,
                position=pos,
                process=None,  # 合并时再标 SMT/DIP
                remark=remark,
                sort_order=sort_i,
            )
        )
    if not lines:
        raise ValueError("亿兰科 BOM 未解析到物料行")

    role_n = (role or "").strip().lower()
    if role_n not in ("smt", "dip"):
        role_n = detect_yilanke_role(rows, header_idx, filename, lines)

    if role_n == "smt":
        model_code = meta.get("smt_semi_code") or Path(filename).stem.upper()
        model_name = f"SMT半成品 {model_code}"
    else:
        model_code = meta.get("finished_code") or Path(filename).stem.upper()
        model_name = f"整机 {model_code}"

    for ln in lines:
        ln.process = "SMT" if role_n == "smt" else "DIP"

    return ParsedBom(
        model_code=model_code,
        model_name=model_name,
        model_spec=meta.get("title"),
        lines=lines,
        folder_name="",
        source_file=filename or "",
        source_mtime=0.0,
    )


def _sheet_score(name: str, header_idx: Optional[int]) -> int:
    """优先「生产用」sheet；有合法表头的 sheet 才计入。"""
    if header_idx is None:
        return -1
    n = (name or "").strip()
    if n == "生产用":
        return 100
    if "生产" in n:
        return 80
    return 50 - header_idx  # 同档次偏表头更靠前的


def _load_yilanke_rows_from_bytes(content: bytes, filename: str) -> tuple[list, int, str]:
    """读全部 sheet，挑含亿兰科表头的最优一份（避免只读 Sheet1 撞上 ERP 导出）。"""
    import tempfile

    suffix = Path(filename or "bom.xlsx").suffix.lower() or ".xlsx"
    with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
        tmp.write(content)
        path = Path(tmp.name)
    try:
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            best: Optional[tuple[int, list, int, str]] = None
            for ws in wb.worksheets:
                rows = list(ws.iter_rows(values_only=True))
                header_idx = find_yilanke_header_row(rows)
                score = _sheet_score(ws.title, header_idx)
                if score < 0 or header_idx is None:
                    continue
                cand = (score, rows, header_idx, ws.title)
                if best is None or cand[0] > best[0]:
                    best = cand
            if best is None:
                # 仍返回首 sheet，便于上层报错时附带预览
                ws0 = wb.worksheets[0]
                rows0 = list(ws0.iter_rows(values_only=True))
                return rows0, -1, ws0.title
            return best[1], best[2], best[3]
        finally:
            wb.close()
    finally:
        try:
            path.unlink(missing_ok=True)
        except OSError:
            pass


def parse_yilanke_bytes(content: bytes, filename: str, *, role: str = "") -> ParsedBom:
    rows, header_idx, sheet_name = _load_yilanke_rows_from_bytes(content, filename)
    if header_idx is None or header_idx < 0:
        raise ValueError(
            f"未找到亿兰科 BOM 表头（子项物料编码 + 用量）。文件：{filename}"
            + (f"（已扫 sheet：{sheet_name}）" if sheet_name else "")
        )
    parsed = parse_yilanke_rows(rows, header_idx, role=role, filename=filename)
    parsed.source_file = filename
    return parsed


def _is_smt_semi_line(ln: ParsedBomLine, smt_model_code: str = "") -> bool:
    code = (ln.material_code or "").upper()
    name = (ln.material_name or "").upper().replace(" ", "")
    if smt_model_code and code == smt_model_code.upper():
        return True
    if code.startswith("3003-"):
        return True
    if "半成品" in (ln.material_name or "") or "SMD半成品" in name:
        return True
    return False


def merge_yilanke_smt_dip(smt: ParsedBom, dip: ParsedBom) -> ParsedBom:
    """
    以 DIP 整机料号为系统机型：去掉 DIP 中的 SMT 半成品行，并入全部 SMT 行。
    """
    smt_code = (smt.model_code or "").upper()
    finished = (dip.model_code or "").upper()
    if not finished:
        # 从 SMT 标题里的括号整机号回填
        m = _RE_FINISHED.search(smt.model_spec or "")
        if m:
            finished = m.group(1).upper()
    if not finished:
        raise ValueError("无法识别整机料号（3001-xxxx），请确认 DIP 文件标题")

    merged: list[ParsedBomLine] = []
    sort_i = 0
    for ln in smt.lines:
        sort_i += 1
        merged.append(
            ParsedBomLine(
                seq=ln.seq,
                material_code=ln.material_code,
                material_name=ln.material_name,
                spec=ln.spec,
                unit=ln.unit,
                qty_per=ln.qty_per,
                position=ln.position,
                process="SMT",
                remark=ln.remark,
                sort_order=sort_i,
            )
        )
    skipped_semi = 0
    for ln in dip.lines:
        if _is_smt_semi_line(ln, smt_code):
            skipped_semi += 1
            continue
        sort_i += 1
        merged.append(
            ParsedBomLine(
                seq=ln.seq,
                material_code=ln.material_code,
                material_name=ln.material_name,
                spec=ln.spec,
                unit=ln.unit,
                qty_per=ln.qty_per,
                position=ln.position,
                process="DIP",
                remark=ln.remark,
                sort_order=sort_i,
            )
        )
    if not merged:
        raise ValueError("合并后 BOM 为空")

    src = f"{smt.source_file}+{dip.source_file}"
    return ParsedBom(
        model_code=finished,
        model_name=dip.model_name or f"整机 {finished}",
        model_spec=f"亿兰科合并 SMT({smt_code or '?'})+DIP；去半成品行{skipped_semi}",
        lines=merged,
        folder_name=finished,
        source_file=src,
        source_mtime=max(float(smt.source_mtime or 0), float(dip.source_mtime or 0)),
    )


def finalize_yilanke_single(parsed: ParsedBom) -> ParsedBom:
    """纯插件 / 纯贴片：仅一份 BOM，不做 SMT+DIP 合并。"""
    role = (parsed.lines[0].process or "DIP").strip().upper() if parsed.lines else "DIP"
    # 以行内工艺为准；若混杂则按多数
    smt_n = sum(1 for ln in parsed.lines if (ln.process or "").upper() == "SMT")
    dip_n = sum(1 for ln in parsed.lines if (ln.process or "").upper() == "DIP")
    if dip_n >= smt_n:
        role = "DIP"
    else:
        role = "SMT"

    model_code = (parsed.model_code or "").strip().upper()
    if not model_code:
        raise ValueError("无法识别机型料号，请确认 Excel 标题或文件名含 3001-/3003-")

    kept: list[ParsedBomLine] = []
    skipped_semi = 0
    sort_i = 0
    for ln in parsed.lines:
        # 纯 DIP 时若仍带 SMT 半成品占位行，去掉（无配套 SMT 表可展开）
        if role == "DIP" and _is_smt_semi_line(ln):
            skipped_semi += 1
            continue
        sort_i += 1
        kept.append(
            ParsedBomLine(
                seq=ln.seq,
                material_code=ln.material_code,
                material_name=ln.material_name,
                spec=ln.spec,
                unit=ln.unit,
                qty_per=ln.qty_per,
                position=ln.position,
                process=role,
                remark=ln.remark,
                sort_order=sort_i,
            )
        )
    if not kept:
        raise ValueError("单份 BOM 解析后无有效物料行")

    if role == "DIP":
        model_name = parsed.model_name or f"整机 {model_code}"
        spec = f"亿兰科纯 DIP（单份）；去半成品行{skipped_semi}"
    else:
        model_name = parsed.model_name or f"SMT半成品 {model_code}"
        spec = "亿兰科纯 SMT（单份）"

    return ParsedBom(
        model_code=model_code,
        model_name=model_name,
        model_spec=spec,
        lines=kept,
        folder_name=model_code,
        source_file=parsed.source_file,
        source_mtime=float(parsed.source_mtime or 0),
    )


def _pair_smt_dip(a: ParsedBom, a_name: str, b: ParsedBom, b_name: str) -> tuple[ParsedBom, ParsedBom]:
    """容错：用户可能对调了 SMT/DIP 文件。"""
    a_is_dip = any(_is_smt_semi_line(ln) for ln in a.lines)
    b_is_dip = any(_is_smt_semi_line(ln) for ln in b.lines)
    if a_is_dip and not b_is_dip:
        return b, a
    if (b.model_code or "").upper().startswith("3003") or (b_name or "").upper().startswith("3003"):
        if not a_is_dip:
            return b, a
        return a, b
    if (a.model_code or "").upper().startswith("3003") or (a_name or "").upper().startswith("3003"):
        return a, b
    # 文件名 3001 → DIP
    if (a_name or "").upper().startswith("3001") and not (b_name or "").upper().startswith("3001"):
        return b, a
    if (b_name or "").upper().startswith("3001") and not (a_name or "").upper().startswith("3001"):
        return a, b
    return a, b


def import_yilanke_pair_bytes(
    db,
    *,
    smt_content: bytes,
    smt_filename: str,
    dip_content: Optional[bytes] = None,
    dip_filename: str = "",
    internal_code: str,
    purchase_no: str,
) -> dict:
    """解析 1～2 份文件并写入 BomModel。两份则合并；一份则按纯 DIP/纯 SMT 入库。"""
    from bom_excel import upsert_bom_model
    from config import get_engineering_customer_by_code

    ic = (internal_code or "").strip().upper()
    pn = (purchase_no or "").strip()
    customer = get_engineering_customer_by_code(ic)
    if not customer:
        return {"status": "failed", "message": f"未知内部代码: {internal_code}", "lines": 0}
    if ic != "A120" and (customer.get("customer_id") or "") != "yilanke":
        return {"status": "failed", "message": "SMT+DIP 合并导入仅支持亿兰科（A120）", "lines": 0}
    if not pn:
        return {"status": "failed", "message": "请先选择在制订单后再导入", "lines": 0}
    if not smt_content:
        return {"status": "failed", "message": "请上传 BOM Excel", "lines": 0}

    single = not dip_content
    if single:
        parsed = parse_yilanke_bytes(smt_content, smt_filename)
        merged = finalize_yilanke_single(parsed)
        smt_lines = sum(1 for ln in merged.lines if (ln.process or "").upper() == "SMT")
        dip_lines = sum(1 for ln in merged.lines if (ln.process or "").upper() == "DIP")
    else:
        a = parse_yilanke_bytes(smt_content, smt_filename)
        b = parse_yilanke_bytes(dip_content, dip_filename or "dip.xlsx")
        smt, dip = _pair_smt_dip(a, smt_filename, b, dip_filename)
        for ln in smt.lines:
            ln.process = "SMT"
        for ln in dip.lines:
            ln.process = "DIP"
        merged = merge_yilanke_smt_dip(smt, dip)
        smt_lines = len(smt.lines)
        dip_lines = len([ln for ln in dip.lines if not _is_smt_semi_line(ln, smt.model_code)])

    customer_id = (customer.get("customer_id") or "").strip()
    customer_name = customer.get("name") or customer_id
    row, created, updated = upsert_bom_model(
        db,
        ic,
        customer_id,
        customer_name,
        merged,
        purchase_no=pn,
    )
    return {
        "status": "success",
        "message": "ok",
        "model_code": row.model_code,
        "lines": len(merged.lines),
        "bom_model_id": row.id,
        "internal_code": ic,
        "purchase_no": pn,
        "file": merged.source_file,
        "_created": created,
        "_updated": updated,
        "smt_lines": smt_lines,
        "dip_lines": dip_lines,
        "single_file": single,
    }
