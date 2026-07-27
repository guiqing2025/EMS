"""PCB 贴片坐标文件解析（Altium Pick&Place / AIS 坐标）"""
from __future__ import annotations

import csv
import re
from dataclasses import dataclass
from pathlib import Path
from typing import Optional


@dataclass
class ParsedPlacementLine:
    refdes: str
    comment: Optional[str] = None
    footprint: Optional[str] = None
    layer: Optional[str] = None
    mid_x: Optional[float] = None
    mid_y: Optional[float] = None
    pad_x: Optional[float] = None
    pad_y: Optional[float] = None
    rotation: Optional[float] = None
    skip: bool = False
    material_hint: Optional[str] = None
    sort_order: int = 0


@dataclass
class ParsedPlacement:
    board_name: str
    file_format: str
    units: str
    lines: list[ParsedPlacementLine]
    source_file: str = ""


def _norm_header(value) -> str:
    if value is None:
        return ""
    # Orient. / Mid X → orient / midx；去掉表头常见标点
    text = re.sub(r"\s+", "", str(value).strip().lower())
    return text.strip(".:：")


def _parse_float(value) -> Optional[float]:
    if value is None or value == "":
        return None
    text = str(value).strip().replace("mm", "").replace("mil", "").replace("th", "")
    try:
        return float(text)
    except (TypeError, ValueError):
        return None


def _pick_col(headers: dict[str, int], *names: str) -> Optional[int]:
    for name in names:
        key = _norm_header(name)
        if key in headers:
            return headers[key]
    return None


def _normalize_layer(value: Optional[str]) -> Optional[str]:
    if not value:
        return None
    text = str(value).strip().upper()
    if text in ("T", "TOP", "TOPLAYER"):
        return "T"
    if text in ("B", "BOTTOM", "BOTTOMLAYER"):
        return "B"
    return None


def _split_altium_txt_tokens(line: str) -> list[str]:
    """按空白切分 Altium TXT 行，保留引号内的 Comment。"""
    tokens: list[str] = []
    i = 0
    n = len(line)
    while i < n:
        while i < n and line[i].isspace():
            i += 1
        if i >= n:
            break
        if line[i] == '"':
            j = i + 1
            while j < n and line[j] != '"':
                j += 1
            tokens.append(line[i + 1 : j].strip())
            i = j + 1 if j < n else n
        else:
            j = i
            while j < n and not line[j].isspace():
                j += 1
            token = line[i:j].strip()
            if token:
                tokens.append(token)
            i = j
    return tokens


def _locate_altium_header_row(rows: list[list[str]], max_rows: int = 80) -> tuple[Optional[int], dict[str, int]]:
    for idx, row in enumerate(rows[:max_rows]):
        mapped = {_norm_header(cell): i for i, cell in enumerate(row) if _norm_header(cell)}
        if _pick_col(mapped, "designator", "位号", "refdes", "refdesignator") is not None:
            return idx, mapped
    return None, {}


def _locate_altium_txt_header(lines_raw: list[str], max_rows: int = 80) -> Optional[int]:
    for idx, line in enumerate(lines_raw[:max_rows]):
        lower = line.lower()
        if "designator" not in lower:
            continue
        if any(
            token in lower
            for token in (
                "mid x",
                "midx",
                "center-x",
                "center-x(mm)",
                "footprint",
                "layer",
                " tb",
                "tb ",
                "rotation",
            )
        ):
            return idx
    return None


def _altium_txt_whitespace_schema(header_line: str) -> dict[str, int]:
    """按 Altium TXT 表头推断 whitespace-split 列序。"""
    hl = re.sub(r"\s+", " ", header_line.strip().lower())
    if re.search(r"designator comment", hl) and "center-x" in hl:
        return {
            "designator": 0,
            "comment": 1,
            "layer": 2,
            "footprint": 3,
            "mid_x": 4,
            "mid_y": 5,
            "rotation": 6,
            "description": 7,
        }
    if re.search(r"designator footprint", hl) and ("mid x" in hl or "midx" in hl):
        return {
            "designator": 0,
            "footprint": 1,
            "mid_x": 2,
            "mid_y": 3,
            "ref_x": 4,
            "ref_y": 5,
            "pad_x": 6,
            "pad_y": 7,
            "layer": 8,
            "rotation": 9,
            "comment": 10,
        }
    if "designator" in hl and ("mid x" in hl or "center-x" in hl):
        return {
            "designator": 0,
            "footprint": 1,
            "mid_x": 2,
            "mid_y": 3,
            "ref_x": 4,
            "ref_y": 5,
            "pad_x": 6,
            "pad_y": 7,
            "layer": 8,
            "rotation": 9,
            "comment": 10,
        }
    return {}


def _line_from_whitespace_schema(
    parts: list[str],
    schema: dict[str, int],
    order: int,
) -> Optional[ParsedPlacementLine]:
    idx_des = schema.get("designator")
    if idx_des is None or idx_des >= len(parts):
        return None
    refdes = parts[idx_des].strip().strip('"')
    if not refdes or refdes.lower() in {"designator", "comment"}:
        return None

    def part(key: str) -> Optional[str]:
        idx = schema.get(key)
        if idx is None or idx >= len(parts):
            return None
        return parts[idx].strip().strip('"')

    return ParsedPlacementLine(
        refdes=refdes,
        comment=part("comment"),
        footprint=part("footprint"),
        layer=_normalize_layer(part("layer")),
        mid_x=_parse_float(part("mid_x")),
        mid_y=_parse_float(part("mid_y")),
        pad_x=_parse_float(part("pad_x")),
        pad_y=_parse_float(part("pad_y")),
        rotation=_parse_float(part("rotation")),
        sort_order=order,
    )


def parse_altium_csv(path: Path) -> ParsedPlacement:
    rows: list[list[str]] = []
    with path.open("r", encoding="utf-8-sig", errors="replace") as fh:
        reader = csv.reader(fh)
        for row in reader:
            rows.append(row)
    header_idx, headers = _locate_altium_header_row(rows)
    if header_idx is None:
        raise ValueError("未找到 Altium CSV 表头（Designator）")

    idx_des = _pick_col(headers, "designator", "位号", "refdes", "refdesignator")
    idx_fp = _pick_col(headers, "footprint", "封装")
    idx_mx = _pick_col(headers, "midx", "mid x", "center-x(mm)", "centerx")
    idx_my = _pick_col(headers, "midy", "mid y", "center-y(mm)", "centery")
    idx_px = _pick_col(headers, "padx", "pad x", "pad-x(mm)")
    idx_py = _pick_col(headers, "pady", "pad y", "pad-y(mm)")
    idx_layer = _pick_col(headers, "layer", "tb")
    idx_rot = _pick_col(headers, "rotation")
    idx_comment = _pick_col(headers, "comment", "值")

    lines: list[ParsedPlacementLine] = []
    order = 0
    for row in rows[header_idx + 1 :]:
        if not row or idx_des is None or idx_des >= len(row):
            continue
        refdes = str(row[idx_des]).strip().strip('"')
        if not refdes:
            continue
        order += 1
        lines.append(
            ParsedPlacementLine(
                refdes=refdes,
                comment=str(row[idx_comment]).strip() if idx_comment is not None and idx_comment < len(row) else None,
                footprint=str(row[idx_fp]).strip() if idx_fp is not None and idx_fp < len(row) else None,
                layer=_normalize_layer(row[idx_layer] if idx_layer is not None and idx_layer < len(row) else None),
                mid_x=_parse_float(row[idx_mx]) if idx_mx is not None and idx_mx < len(row) else None,
                mid_y=_parse_float(row[idx_my]) if idx_my is not None and idx_my < len(row) else None,
                pad_x=_parse_float(row[idx_px]) if idx_px is not None and idx_px < len(row) else None,
                pad_y=_parse_float(row[idx_py]) if idx_py is not None and idx_py < len(row) else None,
                rotation=_parse_float(row[idx_rot]) if idx_rot is not None and idx_rot < len(row) else None,
                sort_order=order,
            )
        )
    if not lines:
        raise ValueError("坐标文件无有效位号行")
    return ParsedPlacement(
        board_name=path.stem.replace("Pick Place for ", "").strip(),
        file_format="altium_csv",
        units="mm",
        lines=lines,
        source_file=str(path),
    )


def parse_altium_txt(path: Path) -> ParsedPlacement:
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    lines_raw = [ln for ln in text.splitlines() if ln.strip()]

    # 少数 Altium 导出虽为 .txt，实为逗号分隔
    for line in lines_raw[:80]:
        if "designator" in line.lower() and "," in line:
            import io

            rows = list(csv.reader(io.StringIO("\n".join(lines_raw))))
            header_idx, headers = _locate_altium_header_row(rows)
            if header_idx is not None:
                idx_des = _pick_col(headers, "designator", "位号", "refdes", "refdesignator")
                idx_fp = _pick_col(headers, "footprint", "封装")
                idx_mx = _pick_col(headers, "midx", "mid x", "center-x(mm)", "centerx", "center-x")
                idx_my = _pick_col(headers, "midy", "mid y", "center-y(mm)", "centery", "center-y")
                idx_layer = _pick_col(headers, "layer", "tb")
                idx_rot = _pick_col(headers, "rotation")
                idx_comment = _pick_col(headers, "comment", "值")
                lines: list[ParsedPlacementLine] = []
                order = 0
                for row in rows[header_idx + 1 :]:
                    if not row or idx_des is None or idx_des >= len(row):
                        continue
                    refdes = str(row[idx_des]).strip().strip('"')
                    if not refdes:
                        continue
                    order += 1
                    lines.append(
                        ParsedPlacementLine(
                            refdes=refdes,
                            comment=str(row[idx_comment]).strip() if idx_comment is not None and idx_comment < len(row) else None,
                            footprint=str(row[idx_fp]).strip() if idx_fp is not None and idx_fp < len(row) else None,
                            layer=_normalize_layer(row[idx_layer] if idx_layer is not None and idx_layer < len(row) else None),
                            mid_x=_parse_float(row[idx_mx]) if idx_mx is not None and idx_mx < len(row) else None,
                            mid_y=_parse_float(row[idx_my]) if idx_my is not None and idx_my < len(row) else None,
                            rotation=_parse_float(row[idx_rot]) if idx_rot is not None and idx_rot < len(row) else None,
                            sort_order=order,
                        )
                    )
                if lines:
                    return ParsedPlacement(
                        board_name=path.stem.replace("Pick Place for ", "").strip(),
                        file_format="altium_txt",
                        units="mm",
                        lines=lines,
                        source_file=str(path),
                    )
            break

    header_idx = _locate_altium_txt_header(lines_raw)
    if header_idx is None:
        raise ValueError("未找到 Altium TXT 表头")

    header_line = lines_raw[header_idx]
    schema = _altium_txt_whitespace_schema(header_line)
    if schema:
        lines: list[ParsedPlacementLine] = []
        order = 0
        for line in lines_raw[header_idx + 1 :]:
            parts = _split_altium_txt_tokens(line)
            if len(parts) < 6:
                continue
            parsed = _line_from_whitespace_schema(parts, schema, order + 1)
            if not parsed:
                continue
            order += 1
            parsed.sort_order = order
            lines.append(parsed)
        if lines:
            return ParsedPlacement(
                board_name=path.stem.replace("Pick Place for ", "").strip(),
                file_format="altium_txt",
                units="mm",
                lines=lines,
                source_file=str(path),
            )

    headers = re.split(r"\s{2,}|\t", header_line.strip())
    if len(headers) < 3:
        headers = header_line.split()

    norm_headers = [_norm_header(h) for h in headers]
    col_map = {name: i for i, name in enumerate(norm_headers)}

    def col(*names: str) -> Optional[int]:
        for name in names:
            key = _norm_header(name)
            if key in col_map:
                return col_map[key]
        return None

    idx_des = col("designator")
    idx_fp = col("footprint")
    idx_mx = col("midx", "mid x", "center-x(mm)", "center-x")
    idx_my = col("midy", "mid y", "center-y(mm)", "center-y")
    idx_layer = col("tb", "layer")
    idx_rot = col("rotation")
    idx_comment = col("comment")

    lines = []
    order = 0
    for line in lines_raw[header_idx + 1 :]:
        parts = re.split(r"\s{2,}|\t", line.strip())
        if len(parts) < 2:
            parts = _split_altium_txt_tokens(line)
        if idx_des is None or idx_des >= len(parts):
            # 列映射失败时再试 whitespace schema
            fallback = _line_from_whitespace_schema(parts, schema or _altium_txt_whitespace_schema(header_line), order + 1)
            if not fallback:
                continue
            order += 1
            fallback.sort_order = order
            lines.append(fallback)
            continue
        refdes = parts[idx_des].strip().strip('"')
        if not refdes:
            continue
        order += 1
        lines.append(
            ParsedPlacementLine(
                refdes=refdes,
                comment=parts[idx_comment].strip() if idx_comment is not None and idx_comment < len(parts) else None,
                footprint=parts[idx_fp].strip() if idx_fp is not None and idx_fp < len(parts) else None,
                layer=_normalize_layer(parts[idx_layer] if idx_layer is not None and idx_layer < len(parts) else None),
                mid_x=_parse_float(parts[idx_mx]) if idx_mx is not None and idx_mx < len(parts) else None,
                mid_y=_parse_float(parts[idx_my]) if idx_my is not None and idx_my < len(parts) else None,
                rotation=_parse_float(parts[idx_rot]) if idx_rot is not None and idx_rot < len(parts) else None,
                sort_order=order,
            )
        )
    if not lines:
        raise ValueError("坐标文件无有效位号行")
    return ParsedPlacement(
        board_name=path.stem.replace("Pick Place for ", "").strip(),
        file_format="altium_txt",
        units="mm",
        lines=lines,
        source_file=str(path),
    )


def _looks_encrypted_bytes(data: bytes) -> Optional[str]:
    """识别常见文档加密/DRM，避免误报成「格式不对」。"""
    head = data[:256]
    if b"E-SafeNet" in head or b"ESAFENET" in head.upper() or (b"LOCK" in head[:64] and b"SafeNet" in head):
        return (
            "文件已被亿赛通/E-SafeNet 加密，无法解析。"
            "请先在客户侧或本机解密后再导入（解密后应用记事本打开能看到 $PART_SECTION_BEGIN$ 或位号坐标）。"
        )
    if head.startswith(b"%PDF") and b"/Encrypt" in data[:4096]:
        return "这是加密 PDF，不是可贴片坐标文件"
    return None


def _ensure_not_encrypted(path: Path) -> None:
    try:
        data = path.read_bytes()[:4096]
    except OSError:
        return
    msg = _looks_encrypted_bytes(data)
    if msg:
        raise ValueError(msg)


def parse_ais_txt(path: Path) -> ParsedPlacement:
    _ensure_not_encrypted(path)
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    lower_head = "\n".join(text.splitlines()[:30]).lower()
    # Expedition 导出日志常被误当成坐标（仅写 Generating ... vb_ais.txt，无位号数据）
    if (
        "$part_section_begin$" not in lower_head
        and (
            "expedition pcb" in lower_head
            or "generating generic ais" in lower_head
            or "process completed successfully" in lower_head
        )
    ):
        raise ValueError(
            "这是 Expedition 导出日志，不是贴片坐标文件。"
            "请向客户索取真正的 AIS 数据（通常含 $PART_SECTION_BEGIN$ 与位号行，"
            "如 vb_ais.txt / *坐标*.txt 中带 TOP/BOTTOM 的那份），不要上传本日志。"
        )
    # 常见：加密文件被当文本读出乱码，也会落到「缺 PART_SECTION」
    if "$part_section_begin$" not in lower_head:
        enc = _looks_encrypted_bytes(path.read_bytes()[:4096])
        if enc:
            raise ValueError(enc)
        # 二进制乱码启发式
        sample = text[:200]
        if sample and sum(1 for ch in sample if ord(ch) < 9 or (13 < ord(ch) < 32)) > 8:
            raise ValueError(
                "坐标文件内容异常（疑似加密或非文本）。"
                "恩玖 AIS 明文应用记事本打开可见 $PART_SECTION_BEGIN$；"
                "若为亿赛通加密请先解密。"
            )
    units = "mm"
    for line in text.splitlines()[:20]:
        upper = line.strip().upper()
        if upper.startswith("UNITS") or upper.startswith("UUNITS"):
            if "TH" in upper or "MIL" in upper:
                units = "th"
            else:
                units = "mm"

    in_section = False
    lines: list[ParsedPlacementLine] = []
    order = 0
    for raw in text.splitlines():
        line = raw.rstrip()
        if "$PART_SECTION_BEGIN$" in line:
            in_section = True
            continue
        if not in_section or not line.strip() or line.startswith("$"):
            continue
        parts = line.split()
        if len(parts) < 6:
            continue
        refdes = parts[0].strip()
        if not refdes or refdes.startswith("$"):
            continue
        material_hint = parts[1] if len(parts) > 1 else None
        rotation = _parse_float(parts[2]) if len(parts) > 2 else None
        mid_x = _parse_float(parts[3]) if len(parts) > 3 else None
        mid_y = _parse_float(parts[4]) if len(parts) > 4 else None
        layer_raw = parts[5] if len(parts) > 5 else None
        # 第 7 列多为 Mirror（BOTTOM 常见 YES），不是「跳过贴装」
        order += 1
        lines.append(
            ParsedPlacementLine(
                refdes=refdes,
                material_hint=material_hint,
                layer=_normalize_layer(layer_raw),
                mid_x=mid_x,
                mid_y=mid_y,
                rotation=rotation,
                skip=False,
                sort_order=order,
            )
        )
    if not lines:
        if "$PART_SECTION_BEGIN$" not in text.upper():
            raise ValueError(
                "不是有效的 AIS 坐标：缺少 $PART_SECTION_BEGIN$ 与位号数据行。"
                "若文件来自 Expedition，请确认上传的是生成的 AIS 坐标，而不是导出过程日志。"
            )
        raise ValueError("AIS 坐标文件无有效位号行")
    return ParsedPlacement(
        board_name=path.stem.replace("坐标文件", "").strip(),
        file_format="ais_txt",
        units=units,
        lines=lines,
        source_file=str(path),
    )


_SIMPLE_REFDES = re.compile(r"^[A-Za-z][A-Za-z0-9_-]*$")
_UNITS_LINE = re.compile(r"^U?UNITS\s*=", re.I)


def _parse_simple_place_data_line(line: str, order: int) -> Optional[ParsedPlacementLine]:
    """位号 X Y 角度 封装（可无层别），常见于 place_txt / 简易坐标导出。"""
    text = line.strip()
    if not text or text.startswith("#") or text.startswith(";") or _UNITS_LINE.match(text):
        return None
    # Designator MidX MidY ... 带表头行跳过
    lower = text.lower()
    if lower.startswith("designator") or lower.startswith("refdes"):
        return None
    parts = text.split()
    if len(parts) < 5:
        return None
    refdes = parts[0].strip()
    if not _SIMPLE_REFDES.match(refdes):
        return None
    mid_x = _parse_float(parts[1])
    mid_y = _parse_float(parts[2])
    rotation = _parse_float(parts[3])
    if mid_x is None or mid_y is None or rotation is None:
        return None
    # 可选：第 5 列为 TOP/BOT，其余为封装；或第 5 列起全是封装
    layer = None
    footprint_parts = parts[4:]
    if footprint_parts:
        maybe_layer = _normalize_layer(footprint_parts[0])
        if maybe_layer and len(footprint_parts) >= 2:
            layer = maybe_layer
            footprint_parts = footprint_parts[1:]
    footprint = " ".join(footprint_parts).strip() or None
    return ParsedPlacementLine(
        refdes=refdes,
        footprint=footprint,
        layer=layer,
        mid_x=mid_x,
        mid_y=mid_y,
        rotation=rotation,
        sort_order=order,
    )


def _looks_like_simple_place_txt(lines_raw: list[str]) -> bool:
    hits = 0
    has_units = False
    for line in lines_raw[:80]:
        if _UNITS_LINE.match(line.strip()):
            has_units = True
            continue
        if _parse_simple_place_data_line(line, 1):
            hits += 1
            if hits >= 3:
                return True
    return has_units and hits >= 1


def parse_simple_place_txt(path: Path) -> ParsedPlacement:
    """解析简易空白分隔坐标：UUNITS/UNITS + 位号 X Y 角度 封装。"""
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    lines_raw = [ln for ln in text.splitlines() if ln.strip()]
    units = "mm"
    for line in lines_raw[:20]:
        upper = line.strip().upper()
        if _UNITS_LINE.match(upper) or upper.startswith("UNITS") or upper.startswith("UUNITS"):
            # MILLIMETERS 含 "MIL" 子串，不可用简单 in 判断
            if "MILLIMETER" in upper or re.search(r"(^|[^A-Z])MM([^A-Z]|$)", upper):
                units = "mm"
            elif re.search(r"(^|[^A-Z])(TH|MIL|MILS)([^A-Z]|$)", upper):
                units = "th"
            else:
                units = "mm"
            break

    lines: list[ParsedPlacementLine] = []
    order = 0
    for raw in lines_raw:
        order += 1
        parsed = _parse_simple_place_data_line(raw, order)
        if parsed:
            lines.append(parsed)
    if not lines:
        raise ValueError("简易坐标文件无有效位号行（需：位号 X Y 角度 封装）")
    # 重排 sort_order 连续
    for i, row in enumerate(lines, start=1):
        row.sort_order = i
    return ParsedPlacement(
        board_name=path.stem.replace("place_txt", "").replace("place", "").strip() or path.stem,
        file_format="simple_place_txt",
        units=units,
        lines=lines,
        source_file=str(path),
    )


def _looks_like_asc_header(line: str) -> bool:
    """永联/Expedition ASC，或 PADS 风格 RefDes + X/Y + Layer/Orient。"""
    lower = line.lower()
    has_des = (
        "refdesignator" in lower
        or "refdes" in lower
        or ("designator" in lower and ("mounttype" in lower or "mid" in lower or "layer" in lower))
        or "位号" in line
    )
    if not has_des:
        return False
    has_xy = bool(re.search(r"\bx\b", lower) and re.search(r"\by\b", lower))
    has_side = "side" in lower or "layer" in lower or "层" in line
    has_cell = "cellname" in lower or "footprint" in lower or "partdecal" in lower or "封装" in line
    has_rot = "rot" in lower or "orient" in lower or "角度" in line
    return has_xy and (has_side or has_cell or has_rot)


def _parse_asc_rows(rows: list[list], *, file_format: str, path: Path) -> ParsedPlacement:
    header_idx = None
    headers: dict[str, int] = {}
    for idx, row in enumerate(rows[:20]):
        mapped = {_norm_header(cell): i for i, cell in enumerate(row) if _norm_header(cell)}
        # 亿兰科等 PADS 导出：RefDes + Orient. + Layer + PartDecal
        if _pick_col(mapped, "refdesignator", "designator", "refdes", "位号") is not None and (
            _pick_col(mapped, "x", "midx", "midx(mm)", "中心x") is not None
            and _pick_col(mapped, "y", "midy", "midy(mm)", "中心y") is not None
        ):
            header_idx = idx
            headers = mapped
            break
    if header_idx is None:
        raise ValueError("未找到 ASC 坐标表头（RefDesignator）")

    idx_des = _pick_col(headers, "refdesignator", "designator", "refdes", "位号")
    idx_fp = _pick_col(headers, "cellname", "footprint", "partdecal", "封装")
    idx_comment = _pick_col(headers, "parttype", "comment", "零件型号", "物料")
    idx_mx = _pick_col(headers, "x", "midx", "mid x", "midx(mm)", "中心x")
    idx_my = _pick_col(headers, "y", "midy", "mid y", "midy(mm)", "中心y")
    idx_layer = _pick_col(headers, "side", "layer", "tb", "层", "层别")
    idx_rot = _pick_col(headers, "rot", "rotation", "orient", "orientation", "角度")

    lines: list[ParsedPlacementLine] = []
    order = 0
    for row in rows[header_idx + 1 :]:
        if not row or idx_des is None or idx_des >= len(row):
            continue
        refdes = str(row[idx_des]).strip().strip('"')
        if not refdes:
            continue
        order += 1
        fp = None
        if idx_fp is not None and idx_fp < len(row) and row[idx_fp] is not None:
            fp = str(row[idx_fp]).strip().strip('"') or None
        comment = None
        if idx_comment is not None and idx_comment < len(row) and row[idx_comment] is not None:
            comment = str(row[idx_comment]).strip().strip('"') or None
        lines.append(
            ParsedPlacementLine(
                refdes=refdes,
                comment=comment,
                footprint=fp,
                layer=_normalize_layer(row[idx_layer] if idx_layer is not None and idx_layer < len(row) else None),
                mid_x=_parse_float(row[idx_mx]) if idx_mx is not None and idx_mx < len(row) else None,
                mid_y=_parse_float(row[idx_my]) if idx_my is not None and idx_my < len(row) else None,
                rotation=_parse_float(row[idx_rot]) if idx_rot is not None and idx_rot < len(row) else None,
                sort_order=order,
            )
        )
    if not lines:
        raise ValueError("ASC 坐标表无有效位号行")
    board_name = path.stem.replace("_ASC", "").replace("_asc", "").strip() or path.stem
    return ParsedPlacement(
        board_name=board_name,
        file_format=file_format,
        units="mm",
        lines=lines,
        source_file=str(path),
    )


def parse_asc_txt(path: Path) -> ParsedPlacement:
    """解析永联 PCBA 包内 *_ASC.txt（空白分列 + RefDesignator）。"""
    text = path.read_text(encoding="utf-8-sig", errors="replace")
    lines_raw = [ln for ln in text.splitlines() if ln.strip()]
    if not lines_raw:
        raise ValueError("ASC 坐标文件为空")
    header_idx = None
    for idx, line in enumerate(lines_raw[:20]):
        if _looks_like_asc_header(line):
            header_idx = idx
            break
    if header_idx is None:
        raise ValueError("未找到 ASC 坐标表头（RefDesignator）")
    rows = [_split_altium_txt_tokens(ln) for ln in lines_raw]
    return _parse_asc_rows(rows, file_format="asc_txt", path=path)


def parse_asc_xlsx(path: Path) -> ParsedPlacement:
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        rows = list(ws.iter_rows(values_only=True))
    finally:
        wb.close()
    if not rows:
        raise ValueError("ASC 坐标表为空")
    try:
        return _parse_asc_rows(rows, file_format="asc_xlsx", path=path)
    except ValueError:
        # 恩玖等常见：无表头，列顺序为 位号 X Y 角度 层别 [料号/封装]
        return _parse_simple_xy_rows(rows, file_format="simple_xy_xlsx", path=path)


def _parse_simple_xy_rows(rows: list, *, file_format: str, path: Path) -> ParsedPlacement:
    """无表头简易坐标：位号, X, Y, 角度, 层别[, 料号/封装]。"""
    lines: list[ParsedPlacementLine] = []
    order = 0
    for row in rows:
        if not row or row[0] is None:
            continue
        refdes = str(row[0]).strip()
        if not refdes or refdes.lower() in {"refdes", "designator", "refdesignator", "位号"}:
            # 若首行其实是表头且能走 ASC，上面已处理；这里跳过表头字样
            if any(
                str(c).strip().lower() in {"x", "y", "rot", "rotation", "layer", "side", "top", "bottom"}
                for c in row[1:5]
                if c is not None
            ):
                continue
            if refdes.lower() in {"refdes", "designator", "refdesignator", "位号"}:
                continue
        if len(row) < 4:
            continue
        mid_x = _parse_float(row[1])
        mid_y = _parse_float(row[2])
        if mid_x is None or mid_y is None:
            continue
        rotation = _parse_float(row[3]) if len(row) > 3 else None
        layer = _normalize_layer(row[4] if len(row) > 4 else None)
        material_hint = None
        footprint = None
        if len(row) > 5 and row[5] is not None:
            hint = str(row[5]).strip()
            if hint:
                # 数字料站码 → material；其余当封装提示
                if re.fullmatch(r"\d{5,}", hint):
                    material_hint = hint
                else:
                    footprint = hint
        order += 1
        lines.append(
            ParsedPlacementLine(
                refdes=refdes,
                material_hint=material_hint,
                footprint=footprint,
                layer=layer,
                mid_x=mid_x,
                mid_y=mid_y,
                rotation=rotation,
                sort_order=order,
            )
        )
    if not lines:
        raise ValueError("简易坐标表无有效位号行（需：位号 X Y 角度 层别）")
    return ParsedPlacement(
        board_name=path.stem.replace("坐标文件", "").strip() or path.stem,
        file_format=file_format,
        units="mm",
        lines=lines,
        source_file=str(path),
    )


def parse_placement_file(path: Path) -> ParsedPlacement:
    _ensure_not_encrypted(path)
    name_lower = path.name.lower()
    suffix = path.suffix.lower()
    if "fiducial" in name_lower or name_lower.endswith("_mark.txt") or "_mark." in name_lower:
        raise ValueError("基准点文件不含贴片位号，请使用 ASC 坐标表或 Pick&Place")
    # xlsx 绝不能按 AIS 文本读（文件名含「坐标」时尤易误入）
    if suffix in (".xlsx", ".xlsm"):
        if "_asc" in name_lower or path.parent.name.upper().endswith("_ASC"):
            return parse_asc_xlsx(path)
        return parse_asc_xlsx(path)
    if suffix == ".txt" and (
        "_asc" in name_lower or path.parent.name.upper().endswith("_ASC")
    ):
        return parse_asc_txt(path)
    if "pick place" in name_lower and suffix == ".csv":
        return parse_altium_csv(path)
    if "pick place" in name_lower and suffix == ".txt":
        return parse_altium_txt(path)
    if suffix == ".txt" and ("坐标" in path.name or "coordinate" in name_lower or "ais" in name_lower):
        return parse_ais_txt(path)
    if suffix == ".csv":
        return parse_altium_csv(path)
    if suffix == ".txt":
        text = path.read_text(encoding="utf-8-sig", errors="replace")
        lines_raw = [ln for ln in text.splitlines() if ln.strip()]
        joined = "\n".join(lines_raw[:80]).lower()
        if any(_looks_like_asc_header(ln) for ln in lines_raw[:20]):
            return parse_asc_txt(path)
        if "$part_section_begin$" in joined:
            return parse_ais_txt(path)
        if _locate_altium_txt_header(lines_raw) is not None or any(
            "designator" in ln.lower() and "," in ln for ln in lines_raw[:40]
        ):
            return parse_altium_txt(path)
        if _looks_like_simple_place_txt(lines_raw) or "place" in name_lower:
            return parse_simple_place_txt(path)
        try:
            return parse_ais_txt(path)
        except ValueError:
            try:
                return parse_altium_txt(path)
            except ValueError:
                return parse_simple_place_txt(path)
    raise ValueError(f"不支持的坐标文件格式: {path.name}")
