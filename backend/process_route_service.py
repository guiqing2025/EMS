"""机型工序对照（工艺明细.xlsx + 手工勾选）"""
from __future__ import annotations

import json
import re
from datetime import datetime
from pathlib import Path
from typing import Optional

from sqlalchemy.orm import Session

from config import get_process_detail_file_path
from engineering_service import normalize_code
from models import BomModel, ModelProcessRoute

PROCESS_STEP_DEFS = [
    {"key": "laser_label", "label": "镭雕和手工贴码（可选）"},
    {"key": "smt", "label": "SMT-AOI"},
    {"key": "pre_oven_aoi", "label": "炉前AOI"},
    {"key": "insert", "label": "插件"},
    {"key": "post_solder", "label": "后焊"},
    {"key": "post_oven_label", "label": "炉后贴码（后焊一枪记插件+后焊）"},
    {"key": "test", "label": "ICT测试"},
    {
        "key": "conformal",
        "label": "三防",
        "sub_options": ["普通三防", "UV胶"],
    },
]

# 布尔工序字段（不含三防子选项）；potting 仅兼容旧数据，界面不再展示
_BOOL_STEP_KEYS = (
    "laser_label",
    "smt",
    "pre_oven_aoi",
    "insert",
    "post_solder",
    "post_oven_label",
    "test",
    "potting",
)


def default_steps() -> dict:
    return {
        "laser_label": False,
        "smt": False,
        "pre_oven_aoi": False,
        "insert": False,
        "post_solder": False,
        "post_oven_label": False,
        "test": False,
        "conformal": {"enabled": False, "type": "普通三防"},
        "potting": False,
    }


def _parse_steps_json(raw: Optional[str]) -> dict:
    if not raw:
        return default_steps()
    try:
        data = json.loads(raw)
        base = default_steps()
        for key in _BOOL_STEP_KEYS:
            if key in data:
                base[key] = bool(data[key])
        if isinstance(data.get("conformal"), dict):
            base["conformal"] = {
                "enabled": bool(data["conformal"].get("enabled")),
                "type": data["conformal"].get("type") or "普通三防",
            }
        # 旧数据无后焊字段：有插件则默认勾后焊（产线通常连着走）
        if "post_solder" not in data and base.get("insert"):
            base["post_solder"] = True
        return base
    except (json.JSONDecodeError, TypeError):
        return default_steps()


def steps_to_json(steps: dict) -> str:
    base = default_steps()
    for key in _BOOL_STEP_KEYS:
        base[key] = bool(steps.get(key))
    # 炉后贴码机型必须有插件+后焊工序
    if base.get("post_oven_label"):
        base["insert"] = True
        base["post_solder"] = True
    conf = steps.get("conformal") or {}
    base["conformal"] = {
        "enabled": bool(conf.get("enabled")),
        "type": conf.get("type") or "普通三防",
    }
    return json.dumps(base, ensure_ascii=False)


def parse_process_text(text: str) -> dict:
    """将工艺明细表中的工艺字符串解析为勾选状态。"""
    steps = default_steps()
    if not text:
        return steps
    raw = str(text).strip()
    upper = raw.upper()

    if re.search(r"镭雕|贴码|手工贴码", raw):
        steps["laser_label"] = True
    if "SMT" in upper:
        steps["smt"] = True
    if re.search(r"炉前\s*AOI|炉前AOI", raw, re.I):
        steps["pre_oven_aoi"] = True
    if "插件" in raw:
        steps["insert"] = True
        # 工艺明细常不写「后焊」，有插件默认带后焊
        steps["post_solder"] = True
    if re.search(r"后焊", raw):
        steps["post_solder"] = True
    if re.search(r"炉后贴码|过炉后贴", raw):
        steps["post_oven_label"] = True
        steps["insert"] = True
        steps["post_solder"] = True
    if re.search(r"ICT|测试", raw, re.I):
        steps["test"] = True
    if "灌胶" in raw:
        steps["potting"] = True
    elif "电子胶" in raw and "透明" not in raw:
        steps["potting"] = True

    if "三防" in raw:
        coating_type = "普通三防"
        if re.search(r"UV", raw, re.I):
            coating_type = "UV胶"
        steps["conformal"] = {"enabled": True, "type": coating_type}

    return steps


def build_route_display(steps: dict) -> str:
    parts: list[str] = []
    if steps.get("laser_label"):
        parts.append("镭雕/贴码")
    if steps.get("smt"):
        parts.append("SMT-AOI")
    if steps.get("pre_oven_aoi"):
        parts.append("炉前AOI")
    if steps.get("insert"):
        parts.append("插件")
    if steps.get("post_solder"):
        parts.append("后焊")
    if steps.get("post_oven_label"):
        parts.append("炉后贴码")
    if steps.get("test"):
        parts.append("ICT测试")
    conf = steps.get("conformal") or {}
    if conf.get("enabled"):
        ctype = conf.get("type") or "普通三防"
        parts.append(f"三防({ctype})" if ctype != "普通三防" else "三防")
    if steps.get("potting"):
        parts.append("灌胶")
    return "-".join(parts)


def is_eng_docs_exempt_steps(steps: dict | None) -> bool:
    """只灌胶（无 SMT/插件/后焊/三防）→ 免 BOM/坐标/Gerber 导入。"""
    s = steps or {}
    if not s.get("potting"):
        return False
    if s.get("smt") or s.get("pre_oven_aoi") or s.get("insert") or s.get("post_solder") or s.get("post_oven_label"):
        return False
    conf = s.get("conformal") or {}
    if isinstance(conf, dict):
        if conf.get("enabled"):
            return False
    elif conf:
        return False
    return True


def eng_docs_exempt_model_norms(db: Session, *, internal_code: str = "") -> set[str]:
    """工序对照判定为免工程资料导入的机型（规范化料号集合）。"""
    ic = (internal_code or "").strip().upper()
    q = db.query(ModelProcessRoute)
    if ic:
        q = q.filter(ModelProcessRoute.internal_code == ic)
    out: set[str] = set()
    for row in q.all():
        if is_eng_docs_exempt_steps(_parse_steps_json(row.steps_json)):
            norm = normalize_code(row.model_code)
            if norm:
                out.add(norm)
    return out


def is_eng_docs_exempt_model(
    db: Session,
    *,
    model_code: str,
    internal_code: str = "",
) -> bool:
    route = find_route_by_model(db, model_code, internal_code=internal_code)
    if not route:
        return False
    return is_eng_docs_exempt_steps(route.get("steps"))


# 扫码卡控：工序对照找不到/未配置时按最严（插件扫码非必须，后焊统一过站）
STRICT_SCAN_GATE_STEPS = {
    "laser_label": True,
    "smt": True,
    "pre_oven_aoi": True,
    "insert": False,
    "post_solder": True,
    "post_oven_label": False,
    "test": True,
    "conformal": {"enabled": True, "type": "普通三防"},
    "potting": False,
}


def _customer_internal_code(customer_id: str = "") -> str:
    cid = (customer_id or "").strip()
    if not cid:
        return ""
    try:
        from config import get_engineering_customer_by_srm_id

        # 含 customer_id_aliases（手工录单如 manual_a58399fe → A114）
        eng = get_engineering_customer_by_srm_id(cid)
        if eng:
            return (eng.get("internal_code") or "").strip().upper()
    except Exception:
        return ""
    return ""


def find_route_by_model(
    db: Session,
    model_code: str,
    *,
    internal_code: str = "",
) -> Optional[dict]:
    """按机型查找工序对照；可带内部代码缩小范围。"""
    code = normalize_code(model_code)
    if not code:
        return None
    ic = (internal_code or "").strip().upper()
    q = db.query(ModelProcessRoute)
    if ic:
        q = q.filter(ModelProcessRoute.internal_code == ic)
    rows = q.all()
    hits = [route_to_dict(r) for r in rows if normalize_code(r.model_code) == code]
    if not hits:
        return None
    if ic:
        return hits[0]
    # 无内部代码：仅当唯一命中时采用，避免同料号跨客户歧义
    if len(hits) == 1:
        return hits[0]
    configured = [h for h in hits if (h.get("status") == "configured" and h.get("route_display"))]
    if len(configured) == 1:
        return configured[0]
    return None


def resolve_scan_gate_steps(
    db: Session,
    *,
    model_code: str = "",
    internal_code: str = "",
    customer_id: str = "",
) -> dict[str, object]:
    """
    扫码用工序勾选。找不到/未配置 → 最严（SMT-AOI + ICT 都卡）。
    保存工序对照后下次扫码立即生效（每次读库）。
    """
    ic = (internal_code or "").strip().upper() or _customer_internal_code(customer_id)
    route = find_route_by_model(db, model_code, internal_code=ic) if (model_code or "").strip() else None
    if not route and (model_code or "").strip() and ic:
        # 内部代码不对时再按机型全局找一次
        route = find_route_by_model(db, model_code, internal_code="")
    if not route or not (route.get("route_display") or "").strip():
        steps = dict(STRICT_SCAN_GATE_STEPS)
        steps["conformal"] = dict(STRICT_SCAN_GATE_STEPS["conformal"])
        return {
            "steps": steps,
            "source": "strict_default",
            "route_display": build_route_display(steps),
            "model_code": (model_code or "").strip(),
            "internal_code": ic,
        }
    steps = route.get("steps") or default_steps()
    return {
        "steps": steps,
        "source": "process_route",
        "route_display": route.get("route_display") or build_route_display(steps),
        "model_code": route.get("model_code") or (model_code or "").strip(),
        "internal_code": route.get("internal_code") or ic,
    }


def _link_bom_model(db: Session, internal_code: str, model_code: str) -> Optional[int]:
    code = normalize_code(model_code)
    row = (
        db.query(BomModel)
        .filter(
            BomModel.internal_code == internal_code.strip().upper(),
            BomModel.is_active.is_(True),
        )
        .all()
    )
    for item in row:
        if normalize_code(item.model_code) == code:
            return item.id
    for item in row:
        mc = normalize_code(item.model_code)
        if mc and (code.startswith(mc) or mc.startswith(code)):
            return item.id
    return None


def route_to_dict(row: ModelProcessRoute) -> dict:
    steps = _parse_steps_json(row.steps_json)
    return {
        "id": row.id,
        "internal_code": row.internal_code,
        "model_code": row.model_code,
        "model_name": row.model_name,
        "bom_model_id": row.bom_model_id,
        "source": row.source,
        "raw_process": row.raw_process,
        "steps": steps,
        "route_display": build_route_display(steps),
        "status": row.status,
        "remark": row.remark,
        "synced_at": row.synced_at,
        "updated_at": row.updated_at,
        "updated_by": row.updated_by,
    }


def get_route(db: Session, internal_code: str, model_code: str) -> Optional[dict]:
    code = normalize_code(model_code)
    ic = internal_code.strip().upper()
    rows = (
        db.query(ModelProcessRoute)
        .filter(ModelProcessRoute.internal_code == ic)
        .all()
    )
    for row in rows:
        if normalize_code(row.model_code) == code:
            return route_to_dict(row)
    return None


def save_route(
    db: Session,
    internal_code: str,
    model_code: str,
    steps: dict,
    *,
    model_name: Optional[str] = None,
    remark: Optional[str] = None,
    operator: Optional[str] = None,
    source: str = "manual",
    raw_process: Optional[str] = None,
) -> dict:
    ic = internal_code.strip().upper()
    code = model_code.strip()
    norm = normalize_code(code)
    row = (
        db.query(ModelProcessRoute)
        .filter(ModelProcessRoute.internal_code == ic)
        .all()
    )
    target = None
    for item in row:
        if normalize_code(item.model_code) == norm:
            target = item
            break
    now = datetime.utcnow()
    bom_id = _link_bom_model(db, ic, code)
    if not target:
        target = ModelProcessRoute(internal_code=ic, model_code=code)
        db.add(target)
    target.model_name = model_name or target.model_name
    target.bom_model_id = bom_id
    target.steps_json = steps_to_json(steps)
    target.source = source
    if raw_process is not None:
        target.raw_process = raw_process
    target.status = "configured" if build_route_display(steps) else "pending"
    target.remark = remark
    target.updated_at = now
    target.updated_by = operator
    if source == "excel" and not target.synced_at:
        target.synced_at = now
    db.flush()
    return route_to_dict(target)


def delete_route(db: Session, route_id: int) -> dict:
    """按 id 删除机型工序对照。"""
    row = db.query(ModelProcessRoute).filter(ModelProcessRoute.id == int(route_id)).first()
    if not row:
        raise ValueError("工序对照不存在")
    info = {
        "id": row.id,
        "internal_code": row.internal_code,
        "model_code": row.model_code,
    }
    db.delete(row)
    db.flush()
    return info


def list_routes(
    db: Session,
    internal_code: str = "",
    keyword: str = "",
    status: str = "",
) -> list[dict]:
    q = db.query(ModelProcessRoute).order_by(
        ModelProcessRoute.internal_code.asc(), ModelProcessRoute.model_code.asc()
    )
    if internal_code:
        q = q.filter(ModelProcessRoute.internal_code == internal_code.strip().upper())
    if keyword:
        like = f"%{keyword.strip()}%"
        q = q.filter(
            (ModelProcessRoute.model_code.like(like))
            | (ModelProcessRoute.model_name.like(like))
            | (ModelProcessRoute.raw_process.like(like))
            | (ModelProcessRoute.remark.like(like))
        )
    if status:
        q = q.filter(ModelProcessRoute.status == status.strip())
    return [route_to_dict(row) for row in q.limit(500).all()]


def sync_from_workbook(db: Session) -> dict:
    from openpyxl import load_workbook

    path = Path(get_process_detail_file_path())
    if not path.is_file():
        return {
            "status": "error",
            "message": f"工艺明细文件不存在: {path}",
            "rows_imported": 0,
            "source_file": str(path),
        }

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[wb.sheetnames[0]]
        imported = 0
        now = datetime.utcnow()
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or len(row) < 3:
                continue
            internal_code = str(row[0] or "").strip().upper()
            model_code = str(row[1] or "").strip()
            process_text = str(row[2] or "").strip()
            if not internal_code or not model_code:
                continue
            steps = parse_process_text(process_text)
            save_route(
                db,
                internal_code,
                model_code,
                steps,
                source="excel",
                raw_process=process_text,
                operator="system",
            )
            row_obj = (
                db.query(ModelProcessRoute)
                .filter(
                    ModelProcessRoute.internal_code == internal_code,
                    ModelProcessRoute.model_code == model_code,
                )
                .first()
            )
            if row_obj:
                row_obj.synced_at = now
            imported += 1
    finally:
        wb.close()

    return {
        "status": "success",
        "message": f"已同步 {imported} 条工艺明细",
        "rows_imported": imported,
        "source_file": str(path),
        "synced_at": now,
    }
