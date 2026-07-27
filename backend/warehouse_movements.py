"""从共享盘进销表解析物料进出账明细（来料 / 按订单发料 / 退料）。"""
from __future__ import annotations

import hashlib
import logging
from datetime import datetime, timedelta
from typing import Optional

from openpyxl.workbook.workbook import Workbook
from sqlalchemy.orm import Session

from models import WarehouseMovement

logger = logging.getLogger(__name__)

MOVEMENT_INBOUND = "inbound"
MOVEMENT_ISSUE = "issue"
MOVEMENT_RETURN = "return"
MOVEMENT_OVERISSUE = "overissue"
MOVEMENT_CUSTOMER_RETURN = "customer_return"
MOVEMENT_SMT_LOSS = "smt_loss"
MOVEMENT_ADJUST = "adjust"

MOVEMENT_LABELS = {
    MOVEMENT_INBOUND: "进账",
    MOVEMENT_ISSUE: "发料",
    MOVEMENT_RETURN: "退料",
    MOVEMENT_OVERISSUE: "超领",
    MOVEMENT_CUSTOMER_RETURN: "退客",
    MOVEMENT_SMT_LOSS: "SMT损耗",
    MOVEMENT_ADJUST: "盘亏超领",
}

ALL_MOVEMENT_TYPES = list(MOVEMENT_LABELS.keys())


def _header_map(header_row: tuple) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        key = str(cell).strip().replace("\n", "")
        if key:
            mapping[key] = idx
    return mapping


def _find_header_row(rows: list[tuple], *need: str) -> tuple[Optional[int], dict[str, int]]:
    for idx, row in enumerate(rows[:12]):
        if not row:
            continue
        mapping = _header_map(row)
        if all(any(name in mapping for name in group) for group in need):
            return idx, mapping
        flat = [str(c).strip() if c is not None else "" for c in row]
        if all(any(name in flat for name in group) for group in need):
            return idx, mapping
    return None, {}


def _cell(row: tuple, mapping: dict[str, int], *names: str):
    for name in names:
        idx = mapping.get(name)
        if idx is not None and idx < len(row):
            val = row[idx]
            if val not in (None, ""):
                return val
    return None


def _float(val) -> Optional[float]:
    if val in (None, ""):
        return None
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def _normalize_date(val) -> str:
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d")
    if isinstance(val, (int, float)) and val > 20000:
        try:
            base = datetime(1899, 12, 30) + timedelta(days=float(val))
            return base.strftime("%Y-%m-%d")
        except (OverflowError, ValueError):
            pass
    text = str(val or "").strip()
    if len(text) >= 10 and text[4] == "-":
        return text[:10]
    return text


def _dedupe_key(*parts) -> str:
    raw = "|".join(str(p or "").strip() for p in parts)
    return hashlib.sha1(raw.encode("utf-8")).hexdigest()


def _process_from_text(text: str) -> Optional[str]:
    t = (text or "").strip().upper()
    if not t:
        return None
    if "SMT" in t and "DIP" in t:
        return "smt+dip"
    if "SMT" in t:
        return "smt"
    if "DIP" in t:
        return "dip"
    return t.lower()[:16]


def parse_inbound_movements(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    header_idx, mapping = _find_header_row(rows, ("物料编号", "物料编码"), ("来料数",))
    if header_idx is None:
        return []
    items: list[dict] = []
    for row_idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        code = str(_cell(row, mapping, "物料编号", "物料编码") or "").strip()
        if not code or code in ("物料编号", "物料编码"):
            continue
        qty = _float(_cell(row, mapping, "来料数"))
        if not qty or qty <= 0:
            continue
        doc_date = _normalize_date(_cell(row, mapping, "日期"))
        ref_no = str(_cell(row, mapping, "入库单号") or "").strip()
        inbound_type = str(_cell(row, mapping, "入库类型") or "").strip()
        remark = str(_cell(row, mapping, "备注") or "").strip()
        if inbound_type and remark:
            remark = f"{inbound_type} · {remark}"
        elif inbound_type:
            remark = inbound_type
        items.append(
            {
                "movement_type": MOVEMENT_INBOUND,
                "material_code": code,
                "material_name": str(_cell(row, mapping, "物料名称") or "").strip(),
                "spec": str(_cell(row, mapping, "规格型号") or "").strip(),
                "unit": "PCS",
                "qty": qty,
                "qty_delta": qty,
                "doc_date": doc_date,
                "ref_no": ref_no,
                "remark": remark,
                "dedupe_key": _dedupe_key("in", code, doc_date, ref_no, qty, row_idx),
            }
        )
    return items


def parse_issue_movements(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    header_idx, mapping = _find_header_row(rows, ("订单号",), ("物料编号", "物料编码"), ("出库数", "发料数"))
    if header_idx is None:
        return []
    items: list[dict] = []
    for row_idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        code = str(_cell(row, mapping, "物料编号", "物料编码") or "").strip()
        if not code or code in ("物料编号", "物料编码"):
            continue
        qty = _float(_cell(row, mapping, "出库数", "发料数"))
        if not qty or qty <= 0:
            continue
        order_no = str(_cell(row, mapping, "订单号") or "").strip()
        product_model = str(_cell(row, mapping, "产品机型", "机型") or "").strip()
        order_qty = _float(_cell(row, mapping, "订单数", "订单数量"))
        process = _process_from_text(str(_cell(row, mapping, "工艺") or ""))
        doc_date = _normalize_date(_cell(row, mapping, "日期"))
        remark = str(_cell(row, mapping, "备注") or "").strip()
        items.append(
            {
                "movement_type": MOVEMENT_ISSUE,
                "material_code": code,
                "material_name": str(_cell(row, mapping, "物料名称") or "").strip(),
                "spec": str(_cell(row, mapping, "规格型号") or "").strip(),
                "unit": "PCS",
                "qty": qty,
                "qty_delta": -qty,
                "order_no": order_no,
                "product_model": product_model,
                "order_qty": order_qty,
                "process": process,
                "doc_date": doc_date,
                "remark": remark,
                "dedupe_key": _dedupe_key("issue", order_no, code, doc_date, qty, row_idx),
            }
        )
    return items


def parse_return_movements(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    header_idx, mapping = _find_header_row(rows, ("订单号",), ("物料编号", "物料编码"), ("退料数量",))
    if header_idx is None:
        return []
    items: list[dict] = []
    for row_idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        code = str(_cell(row, mapping, "物料编号", "物料编码") or "").strip()
        if not code or code in ("物料编号", "物料编码"):
            continue
        qty = _float(_cell(row, mapping, "退料数量"))
        if not qty or qty <= 0:
            continue
        order_no = str(_cell(row, mapping, "订单号") or "").strip()
        product_model = str(_cell(row, mapping, "机型", "产品机型") or "").strip()
        order_qty = _float(_cell(row, mapping, "订单数量", "订单数"))
        doc_date = _normalize_date(_cell(row, mapping, "日期"))
        remark = str(_cell(row, mapping, "备注") or "").strip()
        items.append(
            {
                "movement_type": MOVEMENT_RETURN,
                "material_code": code,
                "material_name": str(_cell(row, mapping, "物料名称") or "").strip(),
                "spec": str(_cell(row, mapping, "规格型号") or "").strip(),
                "unit": str(_cell(row, mapping, "单位") or "PCS").strip() or "PCS",
                "qty": qty,
                "qty_delta": qty,
                "order_no": order_no,
                "product_model": product_model,
                "order_qty": order_qty,
                "doc_date": doc_date,
                "remark": remark,
                "dedupe_key": _dedupe_key("return", order_no, code, doc_date, qty, row_idx),
            }
        )
    return items


def parse_overissue_movements(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    header_idx, mapping = _find_header_row(rows, ("物料编号", "物料编码", "料号"), ("数量", "已发超领数"))
    if header_idx is None:
        return []
    items: list[dict] = []
    for row_idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        code = str(_cell(row, mapping, "物料编号", "物料编码", "料号") or "").strip()
        if not code:
            continue
        qty = _float(_cell(row, mapping, "已发超领数", "数量", "产线超领数"))
        if not qty or qty <= 0:
            continue
        order_no = str(_cell(row, mapping, "订单号", "单号") or "").strip()
        product_model = str(_cell(row, mapping, "机型", "产品机型") or "").strip()
        doc_date = _normalize_date(_cell(row, mapping, "日期", "申请日期"))
        ref_no = str(_cell(row, mapping, "申请单号", "单号") or "").strip()
        remark = str(_cell(row, mapping, "备注") or "").strip()
        items.append(
            {
                "movement_type": MOVEMENT_OVERISSUE,
                "material_code": code,
                "material_name": str(_cell(row, mapping, "物料名称", "名称") or "").strip(),
                "spec": str(_cell(row, mapping, "规格型号", "规格") or "").strip(),
                "unit": str(_cell(row, mapping, "单位") or "PCS").strip() or "PCS",
                "qty": qty,
                "qty_delta": -qty,
                "order_no": order_no,
                "product_model": product_model,
                "doc_date": doc_date,
                "ref_no": ref_no,
                "remark": remark,
                "dedupe_key": _dedupe_key("over", order_no, code, doc_date, qty, row_idx),
            }
        )
    return items


def parse_customer_return_movements(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    header_idx, mapping = _find_header_row(rows, ("物料编号", "物料编码"), ("数量",))
    if header_idx is None:
        return []
    items: list[dict] = []
    for row_idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        code = str(_cell(row, mapping, "物料编号", "物料编码") or "").strip()
        if not code:
            continue
        qty = _float(_cell(row, mapping, "数量"))
        if not qty or qty <= 0:
            continue
        doc_date = _normalize_date(_cell(row, mapping, "日期"))
        ref_no = str(_cell(row, mapping, "单号") or "").strip()
        remark = str(_cell(row, mapping, "备注") or "").strip()
        items.append(
            {
                "movement_type": MOVEMENT_CUSTOMER_RETURN,
                "material_code": code,
                "material_name": str(_cell(row, mapping, "物料名称") or "").strip(),
                "spec": str(_cell(row, mapping, "规格型号") or "").strip(),
                "unit": str(_cell(row, mapping, "单位") or "PCS").strip() or "PCS",
                "qty": qty,
                "qty_delta": -qty,
                "doc_date": doc_date,
                "ref_no": ref_no,
                "remark": remark,
                "dedupe_key": _dedupe_key("cust_ret", code, doc_date, ref_no, qty, row_idx),
            }
        )
    return items


def parse_smt_loss_movements(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    header_idx, mapping = _find_header_row(rows, ("物料编号", "物料编码"), ("损耗数",))
    if header_idx is None:
        return []
    items: list[dict] = []
    for row_idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        code = str(_cell(row, mapping, "物料编号", "物料编码") or "").strip()
        if not code:
            continue
        raw = _float(_cell(row, mapping, "损耗数"))
        if not raw:
            continue
        qty = abs(raw)
        if qty <= 0:
            continue
        doc_date = _normalize_date(_cell(row, mapping, "日期"))
        product_model = str(_cell(row, mapping, "备注") or "").strip()
        items.append(
            {
                "movement_type": MOVEMENT_SMT_LOSS,
                "material_code": code,
                "material_name": str(_cell(row, mapping, "物料名称") or "").strip(),
                "spec": str(_cell(row, mapping, "规格型号") or "").strip(),
                "unit": "PCS",
                "qty": qty,
                "qty_delta": -qty,
                "doc_date": doc_date,
                "product_model": product_model,
                "remark": "SMT损耗",
                "dedupe_key": _dedupe_key("smt_loss", code, doc_date, qty, row_idx),
            }
        )
    return items


def parse_adjust_movements(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    header_idx, mapping = _find_header_row(rows, ("物料编号", "物料编码"), ("盘亏超领数",))
    if header_idx is None:
        return []
    items: list[dict] = []
    for row_idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
        code = str(_cell(row, mapping, "物料编号", "物料编码") or "").strip()
        if not code:
            continue
        qty = _float(_cell(row, mapping, "盘亏超领数"))
        if not qty or qty <= 0:
            continue
        doc_date = _normalize_date(_cell(row, mapping, "日期"))
        items.append(
            {
                "movement_type": MOVEMENT_ADJUST,
                "material_code": code,
                "material_name": str(_cell(row, mapping, "物料名称") or "").strip(),
                "spec": str(_cell(row, mapping, "规格型号") or "").strip(),
                "unit": "PCS",
                "qty": qty,
                "qty_delta": -qty,
                "doc_date": doc_date,
                "dedupe_key": _dedupe_key("adjust", code, doc_date, qty, row_idx),
            }
        )
    return items


def collect_movements_from_workbook(wb: Workbook) -> list[dict]:
    items: list[dict] = []
    sheet_parsers = [
        ("来料数", parse_inbound_movements),
        ("用料数&发料单", parse_issue_movements),
        ("产线退料数", parse_return_movements),
        ("退客数", parse_customer_return_movements),
        ("产线超领数", parse_overissue_movements),
        ("产线申请数", parse_overissue_movements),
        ("SMT损耗", parse_smt_loss_movements),
        ("仓库盘亏超领数", parse_adjust_movements),
    ]
    for sheet_name, parser in sheet_parsers:
        if sheet_name not in wb.sheetnames:
            continue
        try:
            parsed = parser(wb[sheet_name])
            items.extend(parsed)
        except Exception:
            logger.exception("解析 sheet %s 失败", sheet_name)
    return items


def sync_movements_from_workbook(
    db: Session,
    wb: Workbook,
    customer_id: str,
    customer_name: str,
    source_file: str,
) -> dict[str, int]:
    """按客户 workbook 刷新共享盘进出账明细（不影响系统手工流水）。"""
    parsed = collect_movements_from_workbook(wb)
    db.query(WarehouseMovement).filter(
        WarehouseMovement.customer_id == customer_id,
        WarehouseMovement.source == "excel_sync",
        WarehouseMovement.source_file == source_file,
    ).delete(synchronize_session=False)

    inserted = 0
    for item in parsed:
        dedupe = _dedupe_key(customer_id, item["dedupe_key"])
        row = WarehouseMovement(
            customer_id=customer_id,
            customer_name=customer_name,
            movement_type=item["movement_type"],
            material_code=item["material_code"],
            material_name=item.get("material_name") or "",
            spec=item.get("spec") or "",
            unit=item.get("unit") or "PCS",
            qty=float(item["qty"]),
            qty_delta=float(item["qty_delta"]),
            order_no=item.get("order_no") or "",
            product_model=item.get("product_model") or "",
            order_qty=item.get("order_qty"),
            process=item.get("process"),
            doc_date=item.get("doc_date") or "",
            ref_no=item.get("ref_no") or "",
            source="excel_sync",
            source_file=source_file,
            dedupe_key=dedupe,
            remark=item.get("remark") or "",
            operator=None,
        )
        db.add(row)
        inserted += 1

    counts = {t: 0 for t in ALL_MOVEMENT_TYPES}
    for item in parsed:
        counts[item["movement_type"]] = counts.get(item["movement_type"], 0) + 1
    return {"total": inserted, **counts}
