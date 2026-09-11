"""镭雕登记：规则 CRUD、Excel 导入、条码归属匹配。"""
from __future__ import annotations

import re
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from models import AoiBoardResult, LaserBatch
from pcba_barcode import normalize_yymmdd, parse_pcba_barcode, split_model_code

SEQ_RANGE_RE = re.compile(r"(?:(\d{2})/)?(\d{1,5})\s*[-–—~]\s*(\d{1,5})")


def _parse_dates(raw) -> list[str]:
    text = str(raw or "").strip()
    if not text:
        return []
    # 260303\260304 或 260303/260304
    parts = re.split(r"[\\/\n,，;；\s]+", text)
    out = []
    for p in parts:
        p = p.strip()
        if not p:
            continue
        try:
            out.append(normalize_yymmdd(p if not p.isdigit() else int(p) if len(p) <= 6 else p))
        except ValueError:
            # 仅 DD：结合上下文不处理，留给带前缀的流水行
            if re.fullmatch(r"\d{2}", p):
                out.append(p)
            else:
                raise
    return out


def _parse_seq_specs(seq_raw, dates: list[str]) -> list[tuple[str, int, int]]:
    """解析对应流水 → [(laser_date, seq_from, seq_to), ...]"""
    text = str(seq_raw or "").strip()
    if not text:
        raise ValueError("对应流水不能为空")
    matches = list(SEQ_RANGE_RE.finditer(text.replace("\r", "\n")))
    if not matches:
        raise ValueError(f"无法解析流水：{text}")

    full_dates = [d for d in dates if len(d) == 6]
    results: list[tuple[str, int, int]] = []
    for i, m in enumerate(matches):
        day_suffix, a, b = m.group(1), int(m.group(2)), int(m.group(3))
        seq_from, seq_to = (a, b) if a <= b else (b, a)
        laser_date = None
        if day_suffix and full_dates:
            for d in full_dates:
                if d.endswith(day_suffix):
                    laser_date = d
                    break
        if not laser_date:
            if len(matches) == 1 and len(full_dates) == 1:
                laser_date = full_dates[0]
            elif i < len(full_dates):
                laser_date = full_dates[i]
            elif len(full_dates) == 1:
                laser_date = full_dates[0]
        if not laser_date:
            raise ValueError(f"流水段无法对应日期：{m.group(0)}")
        results.append((laser_date, seq_from, seq_to))
    return results


def expand_laser_rows(
    *,
    customer_id: str,
    customer_name: str,
    date_raw,
    model_code: str,
    purchase_no: str,
    order_qty: float,
    seq_raw,
    remark: str = "",
    source: str = "manual",
    created_by: str = "",
) -> list[dict]:
    prefix, mid, ver = split_model_code(model_code)
    model = f"{prefix}-{mid}-{ver}"
    dates = _parse_dates(date_raw)
    if not dates:
        raise ValueError("日期不能为空")
    # 若只有 DD 占位，要求流水里带 DD/
    specs = _parse_seq_specs(seq_raw, dates)
    now = datetime.utcnow()
    rows = []
    for laser_date, seq_from, seq_to in specs:
        rows.append(
            {
                "customer_id": customer_id or "feilisi",
                "customer_name": customer_name or None,
                "laser_date": laser_date,
                "model_code": model,
                "model_mid": mid,
                "model_ver": ver,
                "purchase_no": (purchase_no or "").strip(),
                "order_qty": float(order_qty or 0),
                "seq_from": seq_from,
                "seq_to": seq_to,
                "remark": (remark or "").strip() or None,
                "source": source,
                "created_by": (created_by or "").strip() or None,
                "created_at": now,
                "updated_at": now,
            }
        )
    return rows


def create_laser_batches(db: Session, payload: dict, created_by: str = "") -> list[LaserBatch]:
    rows = expand_laser_rows(
        customer_id=payload.get("customer_id") or "feilisi",
        customer_name=payload.get("customer_name") or "",
        date_raw=payload.get("laser_date") or payload.get("date"),
        model_code=payload.get("model_code") or "",
        purchase_no=payload.get("purchase_no") or "",
        order_qty=payload.get("order_qty") or 0,
        seq_raw=payload.get("seq_range") or payload.get("seq_raw") or "",
        remark=payload.get("remark") or "",
        source=payload.get("source") or "manual",
        created_by=created_by,
    )
    if not rows[0]["purchase_no"]:
        raise ValueError("采购订单号不能为空")
    created = []
    for data in rows:
        # 同键覆盖更新
        existing = (
            db.query(LaserBatch)
            .filter(
                LaserBatch.customer_id == data["customer_id"],
                LaserBatch.laser_date == data["laser_date"],
                LaserBatch.model_mid == data["model_mid"],
                LaserBatch.model_ver == data["model_ver"],
                LaserBatch.purchase_no == data["purchase_no"],
                LaserBatch.seq_from == data["seq_from"],
                LaserBatch.seq_to == data["seq_to"],
            )
            .first()
        )
        if existing:
            existing.order_qty = data["order_qty"]
            existing.remark = data["remark"]
            existing.updated_at = datetime.utcnow()
            existing.source = data["source"]
            created.append(existing)
        else:
            row = LaserBatch(**data)
            db.add(row)
            created.append(row)
    db.flush()
    return created


def delete_laser_batch(db: Session, batch_id: int) -> None:
    row = db.query(LaserBatch).filter(LaserBatch.id == batch_id).first()
    if not row:
        raise ValueError("记录不存在")
    db.delete(row)


def list_laser_batches(
    db: Session,
    *,
    customer_id: str = "",
    purchase_no: str = "",
    model_code: str = "",
    date_from: str = "",
    date_to: str = "",
    limit: int = 500,
) -> list[LaserBatch]:
    q = db.query(LaserBatch)
    if customer_id:
        q = q.filter(LaserBatch.customer_id == customer_id)
    if purchase_no:
        q = q.filter(LaserBatch.purchase_no.like(f"%{purchase_no.strip()}%"))
    if model_code:
        q = q.filter(LaserBatch.model_code.like(f"%{model_code.strip()}%"))
    if date_from:
        q = q.filter(LaserBatch.laser_date >= normalize_yymmdd(date_from.replace("-", "")[-6:]))
    if date_to:
        q = q.filter(LaserBatch.laser_date <= normalize_yymmdd(date_to.replace("-", "")[-6:]))
    return q.order_by(LaserBatch.laser_date.desc(), LaserBatch.id.desc()).limit(limit).all()


def find_laser_batch_for_barcode(db: Session, barcode: str) -> Optional[LaserBatch]:
    parsed = parse_pcba_barcode(barcode)
    if not parsed or parsed.get("seq") is None:
        return None
    # 贴码（前缀+流水）：barcode_prefix(12) + 流水落入区间
    if parsed.get("kind") == "enjiu_manual":
        prefix = (parsed.get("barcode_prefix") or "").strip()
        if not prefix:
            return None
        return (
            db.query(LaserBatch)
            .filter(
                LaserBatch.customer_id.in_(("enjiu", "a116")),
                LaserBatch.barcode_prefix == prefix,
                LaserBatch.seq_from <= parsed["seq"],
                LaserBatch.seq_to >= parsed["seq"],
            )
            .order_by(LaserBatch.id.desc())
            .first()
        )
    return (
        db.query(LaserBatch)
        .filter(
            LaserBatch.model_mid == parsed["model_mid"],
            LaserBatch.model_ver == parsed["model_ver"],
            LaserBatch.laser_date == parsed["laser_date"],
            LaserBatch.seq_from <= parsed["seq"],
            LaserBatch.seq_to >= parsed["seq"],
        )
        .order_by(LaserBatch.id.desc())
        .first()
    )


def _parse_enjiu_print_date(raw) -> str:
    """26.5.9 / 26.4.20 / datetime → YYMMDD；缺省用 000101。"""
    if raw is None or str(raw).strip() == "":
        return "000101"
    if hasattr(raw, "strftime"):
        return raw.strftime("%y%m%d")
    text = str(raw).strip()
    m = re.fullmatch(r"(\d{2})\.(\d{1,2})\.(\d{1,2})", text)
    if m:
        yy, mm, dd = m.group(1), int(m.group(2)), int(m.group(3))
        return f"{yy}{mm:02d}{dd:02d}"
    digits = re.sub(r"\D", "", text)
    if len(digits) >= 6:
        try:
            return normalize_yymmdd(digits[-6:])
        except ValueError:
            pass
    return "000101"


def _parse_enjiu_flow_range(raw) -> Optional[tuple[str, int, int]]:
    """0293841126170001-4000 → (prefix12, seq_from, seq_to)。"""
    text = str(raw or "").strip().replace("–", "-").replace("—", "-").replace("~", "-")
    if not text:
        return None
    m = re.fullmatch(r"(\d{12})(\d{4})\s*-\s*(\d{1,4})", text)
    if not m:
        return None
    prefix, start_s, end_s = m.group(1), m.group(2), m.group(3)
    seq_from, seq_to = int(start_s), int(end_s)
    if seq_from > seq_to:
        seq_from, seq_to = seq_to, seq_from
    return prefix, seq_from, seq_to


def _upsert_enjiu_print_batch(
    db: Session,
    *,
    model_code: str,
    purchase_no: str,
    order_qty: float,
    prefix: str,
    seq_from: int,
    seq_to: int,
    version: str,
    laser_date: str,
    remark: str,
    created_by: str,
) -> LaserBatch:
    now = datetime.utcnow()
    ver = (version or "").strip()[:8] or "-"
    mid = prefix[:6]
    existing = (
        db.query(LaserBatch)
        .filter(
            LaserBatch.customer_id == "enjiu",
            LaserBatch.barcode_prefix == prefix,
            LaserBatch.purchase_no == purchase_no,
            LaserBatch.seq_from == seq_from,
            LaserBatch.seq_to == seq_to,
        )
        .first()
    )
    if existing:
        existing.model_code = model_code
        existing.model_mid = mid
        existing.model_ver = ver
        existing.order_qty = order_qty
        existing.laser_date = laser_date
        existing.remark = remark or existing.remark
        existing.source = "print_barcode_xlsx"
        existing.updated_at = now
        return existing
    row = LaserBatch(
        customer_id="enjiu",
        customer_name="客户B",
        laser_date=laser_date,
        model_code=model_code,
        model_mid=mid,
        model_ver=ver,
        purchase_no=purchase_no,
        order_qty=order_qty,
        seq_from=seq_from,
        seq_to=seq_to,
        barcode_prefix=prefix,
        remark=remark or None,
        source="print_barcode_xlsx",
        created_by=created_by or "excel",
        created_at=now,
        updated_at=now,
    )
    db.add(row)
    return row


def resolve_print_barcode_register_path() -> Path:
    from config import get_engineering_share_base, get_laser_print_register_path

    configured = get_laser_print_register_path()
    if configured:
        p = Path(configured)
        if p.is_file():
            return p
    base = (get_engineering_share_base() or "").strip()
    if base:
        # A-生产… 的上一级为「共享-测试软件资料」
        parent = Path(base).resolve().parent
        candidate = parent / "I-表格类" / "打印条码登记表.xlsx"
        if candidate.is_file():
            return candidate
    raise FileNotFoundError("未找到打印条码登记表.xlsx（请检查 LASER_PRINT_REGISTER_PATH / 共享盘 I-表格类）")


def _parse_feilisi_d0_flow_range(
    raw,
    *,
    inherit: Optional[tuple[str, str, str]] = None,
) -> Optional[tuple[str, str, str, int, int]]:
    """客户A打印登记流水 → (model_mid, model_ver, laser_date, seq_from, seq_to)。

    主形式：D02002700726011780001-85500 / DS5000170225122900001-00500
    补打仅流水：81214-81700（需 inherit 主段的 mid/ver/date）
    """
    text = str(raw or "").strip().replace("–", "-").replace("—", "-").replace("~", "-")
    if not text:
        return None
    m = re.fullmatch(
        r"(D[S0][0-9A-Za-z]{6}\d{2})(\d{6})(\d{1,6})\s*-\s*(\d{1,6})",
        text,
        re.I,
    )
    if m:
        head, date_s, start_s, end_s = m.group(1), m.group(2), m.group(3), m.group(4)
        head_u = head.upper()
        # D0/DS + mid6 + ver2
        mid = head_u[2:8]
        ver = head_u[8:10]
        seq_from, seq_to = int(start_s), int(end_s)
        # Excel 偶发 6 位流水：压到 5 位（取末 5 位）以贴合条码规则
        if seq_from > 99999:
            seq_from = int(str(seq_from)[-5:])
        if seq_to > 99999:
            seq_to = int(str(seq_to)[-5:])
        if seq_from > seq_to:
            seq_from, seq_to = seq_to, seq_from
        try:
            laser_date = normalize_yymmdd(date_s)
        except ValueError:
            return None
        return mid, ver, laser_date, seq_from, seq_to
    m2 = re.fullmatch(r"(\d{1,5})\s*-\s*(\d{1,5})", text)
    if m2 and inherit:
        mid, ver, laser_date = inherit
        seq_from, seq_to = int(m2.group(1)), int(m2.group(2))
        if seq_from > seq_to:
            seq_from, seq_to = seq_to, seq_from
        return mid, ver, laser_date, seq_from, seq_to
    return None


def _upsert_feilisi_print_batch(
    db: Session,
    *,
    model_code: str,
    purchase_no: str,
    order_qty: float,
    mid: str,
    ver: str,
    laser_date: str,
    seq_from: int,
    seq_to: int,
    remark: str,
    created_by: str,
) -> LaserBatch:
    now = datetime.utcnow()
    prefix, mid_n, ver_n = split_model_code(model_code)
    # 以条码解析出的 mid/ver 为准（与 ICT 一致）；机型展示仍用表内料号
    use_mid = (mid or mid_n).upper()
    use_ver = (ver or ver_n).zfill(2)
    existing = (
        db.query(LaserBatch)
        .filter(
            LaserBatch.customer_id == "feilisi",
            LaserBatch.laser_date == laser_date,
            LaserBatch.model_mid == use_mid,
            LaserBatch.model_ver == use_ver,
            LaserBatch.purchase_no == purchase_no,
            LaserBatch.seq_from == seq_from,
            LaserBatch.seq_to == seq_to,
        )
        .first()
    )
    if existing:
        existing.model_code = model_code
        existing.order_qty = order_qty
        existing.remark = remark or existing.remark
        existing.source = "print_barcode_xlsx"
        existing.updated_at = now
        return existing
    row = LaserBatch(
        customer_id="feilisi",
        customer_name="客户A",
        laser_date=laser_date,
        model_code=model_code,
        model_mid=use_mid,
        model_ver=use_ver,
        purchase_no=purchase_no,
        order_qty=order_qty,
        seq_from=seq_from,
        seq_to=seq_to,
        barcode_prefix=None,
        remark=remark or None,
        source="print_barcode_xlsx",
        created_by=created_by or "excel",
        created_at=now,
        updated_at=now,
    )
    db.add(row)
    return row


def import_feilisi_print_barcode_register(
    db: Session,
    path: str | Path | None = None,
    *,
    sheet_name: str = "菲利斯",  # 共享盘工作表原名
    created_by: str = "excel",
) -> dict[str, Any]:
    """导入共享盘「打印条码登记表」对应 sheet → laser_batches。"""
    xlsx = Path(path) if path else resolve_print_barcode_register_path()
    wb = load_workbook(xlsx, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"工作表不存在：{sheet_name}，可选 {wb.sheetnames}")
    ws = wb[sheet_name]
    created = updated = skipped = 0
    errors: list[str] = []
    header_row = None
    headers: list[str] = []
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True), start=1):
        vals = [str(c or "").strip() for c in row]
        if "机型" in vals and "订单号" in vals and any("流水" in v for v in vals):
            header_row = i
            headers = vals
            break
    if not header_row:
        raise ValueError("未找到表头（需含 机型/订单号/流水范围）")

    def col(*names: str) -> Optional[int]:
        for n in names:
            if n in headers:
                return headers.index(n)
        return None

    i_model = col("机型")
    i_po = col("订单号")
    i_qty = col("订单数", "订单数量", "数量")
    i_flow = col("流水范围", "流水")
    i_date = col("打印日期")
    i_remark = col("备注")
    i_patch = col("补")
    date_idxs = [i for i, h in enumerate(headers) if h == "打印日期"]
    i_patch_date = date_idxs[1] if len(date_idxs) >= 2 else None
    if None in (i_model, i_po, i_flow):
        raise ValueError(f"表头缺少必要列：{headers}")

    for ridx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        model = str(row[i_model] or "").strip()
        purchase_no = str(row[i_po] or "").strip()
        if (
            not model
            or not purchase_no
            or model == "机型"
            or "xxx" in model.lower()
            or "xxxxxx" in purchase_no.lower()
        ):
            skipped += 1
            continue
        try:
            qty = float(row[i_qty] or 0) if i_qty is not None else 0
            remark = str(row[i_remark] or "").strip() if i_remark is not None else ""
            main = _parse_feilisi_d0_flow_range(row[i_flow])
            if not main:
                raise ValueError(f"无法解析流水范围：{row[i_flow]}")
            ranges: list[tuple[str, str, str, int, int, str]] = []
            mid, ver, laser_date, seq_from, seq_to = main
            ranges.append((mid, ver, laser_date, seq_from, seq_to, remark))
            inherit = (mid, ver, laser_date)
            if i_patch is not None and row[i_patch]:
                patch = _parse_feilisi_d0_flow_range(row[i_patch], inherit=inherit)
                if patch:
                    pm, pv, pd, pf, pt = patch
                    # 补打若自带完整 D0，日期以条码内为准；仅流水则沿用主段日期
                    tag = (remark + " 补打").strip() if remark else "补打"
                    ranges.append((pm, pv, pd, pf, pt, tag))
            for mid_i, ver_i, date_i, sf, st, tag in ranges:
                before = (
                    db.query(LaserBatch.id)
                    .filter(
                        LaserBatch.customer_id == "feilisi",
                        LaserBatch.laser_date == date_i,
                        LaserBatch.model_mid == mid_i,
                        LaserBatch.model_ver == ver_i,
                        LaserBatch.purchase_no == purchase_no,
                        LaserBatch.seq_from == sf,
                        LaserBatch.seq_to == st,
                    )
                    .first()
                )
                _upsert_feilisi_print_batch(
                    db,
                    model_code=model,
                    purchase_no=purchase_no,
                    order_qty=qty,
                    mid=mid_i,
                    ver=ver_i,
                    laser_date=date_i,
                    seq_from=sf,
                    seq_to=st,
                    remark=tag,
                    created_by=created_by,
                )
                if before:
                    updated += 1
                else:
                    created += 1
        except Exception as e:
            errors.append(f"第{ridx}行: {e}")
            if len(errors) >= 40:
                break
    db.flush()
    return {
        "path": str(xlsx),
        "sheet": sheet_name,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
    }


def import_enjiu_print_barcode_register(
    db: Session,
    path: str | Path | None = None,
    *,
    sheet_name: str = "恩玖-鼎雄",  # 共享盘工作表原名
    created_by: str = "excel",
) -> dict[str, Any]:
    """导入共享盘「打印条码登记表」对应 sheet → laser_batches（含 barcode_prefix）。"""
    xlsx = Path(path) if path else resolve_print_barcode_register_path()
    wb = load_workbook(xlsx, data_only=True)
    if sheet_name not in wb.sheetnames:
        raise ValueError(f"工作表不存在：{sheet_name}，可选 {wb.sheetnames}")
    ws = wb[sheet_name]
    created = updated = skipped = 0
    errors: list[str] = []
    # 找表头行
    header_row = None
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=5, values_only=True), start=1):
        vals = [str(c or "").strip() for c in row]
        if "机型" in vals and "订单号" in vals and any("流水" in v for v in vals):
            header_row = i
            headers = vals
            break
    if not header_row:
        raise ValueError("未找到表头（需含 机型/订单号/流水范围）")

    def col(*names: str) -> Optional[int]:
        for n in names:
            if n in headers:
                return headers.index(n)
        return None

    i_model = col("机型")
    i_po = col("订单号")
    i_qty = col("订单数", "订单数量", "数量")
    i_flow = col("流水范围", "流水")
    i_ver = col("版本")
    i_date = col("打印日期")
    i_patch = col("补")
    i_patch_date = None
    # 第二组「打印日期」在补打列后
    date_idxs = [i for i, h in enumerate(headers) if h == "打印日期"]
    if len(date_idxs) >= 2:
        i_patch_date = date_idxs[1]
    if None in (i_model, i_po, i_flow):
        raise ValueError(f"表头缺少必要列：{headers}")

    for ridx, row in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        model = str(row[i_model] or "").strip()
        purchase_no = str(row[i_po] or "").strip()
        if not model or not purchase_no or model == "机型" or "xxx" in purchase_no.lower() or purchase_no.endswith("xxxxxx"):
            skipped += 1
            continue
        try:
            qty = float(row[i_qty] or 0) if i_qty is not None else 0
            ver = str(row[i_ver] or "").strip() if i_ver is not None else ""
            laser_date = _parse_enjiu_print_date(row[i_date] if i_date is not None else None)
            ranges: list[tuple[str, int, int, str, str]] = []
            main = _parse_enjiu_flow_range(row[i_flow])
            if main:
                ranges.append((*main, laser_date, ""))
            else:
                raise ValueError(f"无法解析流水范围：{row[i_flow]}")
            if i_patch is not None and row[i_patch]:
                patch = _parse_enjiu_flow_range(row[i_patch])
                if patch:
                    pd = (
                        _parse_enjiu_print_date(row[i_patch_date])
                        if i_patch_date is not None
                        else laser_date
                    )
                    ranges.append((*patch, pd, "补打"))
            for prefix, seq_from, seq_to, d, tag in ranges:
                before = (
                    db.query(LaserBatch.id)
                    .filter(
                        LaserBatch.customer_id == "enjiu",
                        LaserBatch.barcode_prefix == prefix,
                        LaserBatch.purchase_no == purchase_no,
                        LaserBatch.seq_from == seq_from,
                        LaserBatch.seq_to == seq_to,
                    )
                    .first()
                )
                _upsert_enjiu_print_batch(
                    db,
                    model_code=model,
                    purchase_no=purchase_no,
                    order_qty=qty,
                    prefix=prefix,
                    seq_from=seq_from,
                    seq_to=seq_to,
                    version=ver,
                    laser_date=d,
                    remark=tag,
                    created_by=created_by,
                )
                if before:
                    updated += 1
                else:
                    created += 1
        except Exception as e:
            errors.append(f"第{ridx}行: {e}")
            if len(errors) >= 40:
                break
    db.flush()
    return {
        "path": str(xlsx),
        "sheet": sheet_name,
        "created": created,
        "updated": updated,
        "skipped": skipped,
        "errors": errors,
    }


def rematch_aoi_to_laser(db: Session, limit: int = 5000) -> dict[str, int]:
    """只重挂尚未归属采购单的 AOI 条码，避免每轮扫全表。"""
    rows = (
        db.query(AoiBoardResult)
        .filter((AoiBoardResult.purchase_no.is_(None)) | (AoiBoardResult.purchase_no == ""))
        .order_by(AoiBoardResult.id.desc())
        .limit(limit)
        .all()
    )
    matched = unmatched = 0
    for row in rows:
        batch = find_laser_batch_for_barcode(db, row.barcode)
        if batch:
            row.purchase_no = batch.purchase_no
            row.customer_id = batch.customer_id
            row.laser_batch_id = batch.id
            row.model_code = batch.model_code
            matched += 1
        else:
            unmatched += 1
    return {"scanned": len(rows), "matched": matched, "unmatched": unmatched}


def import_laser_excel(
    db: Session,
    path: str | Path,
    *,
    customer_id: str = "feilisi",
    customer_name: str = "客户A",
    created_by: str = "excel",
) -> dict[str, Any]:
    wb = load_workbook(path, data_only=True)
    ws = wb.active
    headers = [str(c.value or "").strip() for c in next(ws.iter_rows(min_row=1, max_row=1))]
    # 期望：入系统 日期 机型 订单号 订单数 对应流水
    created = 0
    skipped = 0
    errors: list[str] = []
    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue
        try:
            date_raw = row[1]
            model = str(row[2] or "").strip()
            purchase_no = str(row[3] or "").strip()
            qty = float(row[4] or 0)
            seq_raw = row[5]
            if not model or not purchase_no:
                skipped += 1
                continue
            create_laser_batches(
                db,
                {
                    "customer_id": customer_id,
                    "customer_name": customer_name,
                    "laser_date": date_raw,
                    "model_code": model,
                    "purchase_no": purchase_no,
                    "order_qty": qty,
                    "seq_range": seq_raw,
                    "source": "excel",
                },
                created_by=created_by,
            )
            created += 1
        except Exception as e:
            errors.append(f"第{i}行: {e}")
            if len(errors) >= 30:
                break
    return {
        "sheet": ws.title,
        "headers": headers,
        "created_rows": created,
        "skipped": skipped,
        "errors": errors,
    }


def boards_for_purchase(
    db: Session,
    purchase_no: str,
    *,
    customer_id: str = "",
    keyword: str = "",
    result: str = "",
    include_items: bool = True,
    limit: int = 5000,
) -> dict[str, Any]:
    pn = (purchase_no or "").strip()
    from device_board_archive import purchase_board_stats

    stats = purchase_board_stats(db, pn, kind="aoi", customer_id=customer_id)
    total_n = stats["total"]
    pass_n = stats["pass_count"]
    fail_n = stats["fail_count"]

    base = db.query(AoiBoardResult).filter(AoiBoardResult.purchase_no == pn)
    if customer_id:
        base = base.filter(AoiBoardResult.customer_id == customer_id)

    # 明细仍以热表为主；归档通过合计已计入 total/pass/fail
    from sqlalchemy import func

    result_u = func.upper(AoiBoardResult.result)

    items: list[dict] = []
    items_truncated = False
    result_filter = (result or "").strip().upper()
    kw = (keyword or "").strip()
    if include_items and (kw or result_filter):
        q = base
        if kw:
            q = q.filter(AoiBoardResult.barcode.like(f"%{kw}%"))
        if result_filter in ("FAIL", "FALL", "NG"):
            q = q.filter(result_u.in_(("FAIL", "FALL", "NG")))
        elif result_filter == "PASS":
            q = q.filter(result_u == "PASS")
        elif result_filter and result_filter != "ALL":
            q = q.filter(result_u == result_filter)
        item_total = q.count()
        rows = q.order_by(AoiBoardResult.tested_at.desc(), AoiBoardResult.id.desc()).limit(limit).all()
        items_truncated = item_total > limit
        items = [
            {
                "barcode": r.barcode,
                "model_code": r.model_code,
                "laser_date": r.laser_date,
                "seq": r.seq,
                "side": r.side,
                "result": ("FAIL" if (r.result or "").upper() == "FALL" else r.result),
                "machine": r.machine,
                "tested_at": r.tested_at.isoformat() if r.tested_at else None,
                "source_file": r.source_file,
            }
            for r in rows
        ]

    batches = (
        db.query(LaserBatch)
        .filter(LaserBatch.purchase_no == pn)
        .order_by(LaserBatch.laser_date.desc())
        .all()
    )
    if customer_id:
        batches = [b for b in batches if b.customer_id == customer_id]
    return {
        "purchase_no": pn,
        "total": total_n,
        "pass_count": pass_n,
        "fail_count": fail_n,
        "unknown_count": max(0, total_n - pass_n - fail_n),
        "items_limit": limit,
        "items_truncated": items_truncated,
        "result_filter": result_filter or "",
        "laser_batches": [
            {
                "id": b.id,
                "laser_date": b.laser_date,
                "model_code": b.model_code,
                "seq_from": b.seq_from,
                "seq_to": b.seq_to,
                "order_qty": b.order_qty,
            }
            for b in batches
        ],
        "items": items,
    }


def resolve_packing_line_by_laser(db, barcode: str = "", **_kwargs):
    """开发补齐：按镭雕登记尝试匹配订单/批次。"""
    batch = find_laser_batch_for_barcode(db, barcode or "")
    if not batch:
        return None, None, "未匹配到镭雕批次"
    return getattr(batch, "purchase_no", None) or getattr(batch, "order_no", None), batch, None


def filter_orders_by_laser_register(db, orders, *, preserve_keyword: str = ""):
    """开发补齐：无额外过滤，原样返回。"""
    _ = (db, preserve_keyword)
    return list(orders or [])
