"""报价模块 BOM 解析（独立，不写工程 BomModel）"""
from __future__ import annotations

import io
import re
from pathlib import Path
from typing import BinaryIO, Optional, Union

from quote_engine import BomInputLine


def _norm_header(value) -> str:
    if value is None:
        return ""
    return re.sub(r"\s+", "", str(value).strip())


def _cell_str(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def _cell_float(value, default: float = 0.0) -> float:
    if value is None or value == "":
        return default
    try:
        return float(value)
    except (TypeError, ValueError):
        return default


def _pick(headers: dict[str, int], *names: str) -> Optional[int]:
    for name in names:
        key = _norm_header(name)
        if key in headers:
            return headers[key]
    return None


def _rows_from_openpyxl(data: bytes) -> list[list]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), data_only=True, read_only=True)
    # 优先名为 BOM 的 sheet，否则第一张有「元件料号/料号」表头的
    sheets = list(wb.worksheets)
    ordered = sorted(sheets, key=lambda ws: 0 if "bom" in (ws.title or "").lower() else 1)
    for ws in ordered:
        rows = [list(r) for r in ws.iter_rows(values_only=True)]
        if _find_header_row(rows) is not None:
            return rows
    return [list(r) for r in sheets[0].iter_rows(values_only=True)] if sheets else []


def _rows_from_xlrd(data: bytes) -> list[list]:
    import xlrd

    book = xlrd.open_workbook(file_contents=data)
    sheets = list(book.sheets())
    ordered = sorted(sheets, key=lambda sh: 0 if "bom" in (sh.name or "").lower() else 1)
    for sh in ordered:
        rows = [sh.row_values(i) for i in range(sh.nrows)]
        if _find_header_row(rows) is not None:
            return rows
    sh = sheets[0]
    return [sh.row_values(i) for i in range(sh.nrows)]


def _find_header_row(rows: list[list]) -> Optional[int]:
    for i, row in enumerate(rows[:30]):
        headers = {_norm_header(c): idx for idx, c in enumerate(row) if _norm_header(c)}
        if _pick(headers, "元件料号", "物料编码", "料号", "物料料号", "材料编码") is not None:
            return i
        if _pick(headers, "元件品名", "物料名称", "品名") is not None and _pick(
            headers, "单位用量", "用量", "单耗", "数量"
        ) is not None:
            return i
    return None


def parse_quote_bom_bytes(data: bytes, filename: str = "") -> tuple[list[BomInputLine], dict]:
    """解析上传的 BOM，返回行列表与头信息提示。"""
    name = (filename or "").lower()
    if name.endswith(".xls") and not name.endswith(".xlsx"):
        rows = _rows_from_xlrd(data)
    else:
        try:
            rows = _rows_from_openpyxl(data)
        except Exception:
            rows = _rows_from_xlrd(data)

    header_idx = _find_header_row(rows)
    if header_idx is None:
        raise ValueError("未识别到 BOM 表头（需要元件料号/料号等列）")

    header_row = rows[header_idx]
    headers = {_norm_header(c): idx for idx, c in enumerate(header_row) if _norm_header(c)}

    col_code = _pick(headers, "元件料号", "物料编码", "料号", "物料料号", "材料编码")
    col_name = _pick(headers, "元件品名", "物料名称", "品名", "材料名称")
    col_spec = _pick(headers, "元件规格", "规格", "规格型号", "材料规格")
    col_qty = _pick(headers, "单位用量", "用量", "单耗", "数量", "用量PCS")
    col_unit = _pick(headers, "单位")
    col_seq = _pick(headers, "序号", "项次")
    col_pos = _pick(headers, "插件位置", "位号", "位置", "RefDes")
    col_ic = _pick(headers, "IC金额", "P", "IC费", "贴片IC")
    col_dip = _pick(headers, "DIP", "Q", "插件点数", "DIP点数")

    # 鼎雄报价 BOM：表头行里 N/O 可能是数字合计，P 列常无标题——取 DIP 前一列数字列当 P
    if col_dip is not None and col_ic is None:
        # 找 DIP 左侧最近的空标题列或无名列
        for idx in range(col_dip - 1, max(col_dip - 4, -1), -1):
            title = _norm_header(header_row[idx]) if idx < len(header_row) else ""
            if title in ("", "P") or (title.replace(".", "", 1).isdigit() is False and title not in headers):
                # 若该列表体多为小数金额，视为 P
                sample = []
                for r in rows[header_idx + 1 : header_idx + 40]:
                    if idx < len(r):
                        v = r[idx]
                        if isinstance(v, (int, float)) and float(v) != 0:
                            sample.append(float(v))
                if sample and max(sample) < 50:  # IC 金额通常较小
                    col_ic = idx
                    break

    # 表头本身把 304/119 写在 N/O 列名位置时，不要当数据列

    meta = {
        "header_row": header_idx + 1,
        "columns": {
            "code": col_code,
            "name": col_name,
            "spec": col_spec,
            "qty": col_qty,
            "ic": col_ic,
            "dip": col_dip,
        },
    }

    # 尝试从头几行读产品名
    product_code = ""
    product_name = ""
    for r in rows[: header_idx + 1]:
        for c in r:
            t = _cell_str(c)
            if re.match(r"^\d{2,3}-\d+", t):
                product_code = product_code or t
            if "PCBA" in t.upper() or "功率板" in t:
                product_name = product_name or t

    # 主件料号列
    col_main = _pick(headers, "主件料号", "成品料号", "机型料号")
    col_main_name = _pick(headers, "主件品名", "成品名称")

    out: list[BomInputLine] = []
    for raw in rows[header_idx + 1 :]:
        if not raw:
            continue
        code = _cell_str(raw[col_code]) if col_code is not None and col_code < len(raw) else ""
        qty = _cell_float(raw[col_qty]) if col_qty is not None and col_qty < len(raw) else 0
        name = _cell_str(raw[col_name]) if col_name is not None and col_name < len(raw) else ""
        spec = _cell_str(raw[col_spec]) if col_spec is not None and col_spec < len(raw) else ""
        if not code and not name:
            continue
        if qty <= 0 and not name:
            continue
        if col_main is not None and col_main < len(raw):
            mc = _cell_str(raw[col_main])
            if mc and not product_code:
                product_code = mc
        if col_main_name is not None and col_main_name < len(raw):
            mn = _cell_str(raw[col_main_name])
            if mn and not product_name:
                product_name = mn

        ic_amt = _cell_float(raw[col_ic]) if col_ic is not None and col_ic < len(raw) else 0
        dip_pts = _cell_float(raw[col_dip]) if col_dip is not None and col_dip < len(raw) else 0
        # 跳过 PCB 本身可保留给引擎 skip

        out.append(
            BomInputLine(
                seq=_cell_str(raw[col_seq]) if col_seq is not None and col_seq < len(raw) else "",
                material_code=code,
                material_name=name,
                spec=spec,
                qty_per=qty if qty > 0 else 1.0,
                unit=_cell_str(raw[col_unit]) if col_unit is not None and col_unit < len(raw) else "PCS",
                position=_cell_str(raw[col_pos]) if col_pos is not None and col_pos < len(raw) else "",
                ic_amount=ic_amt,
                dip_points=dip_pts,
            )
        )

    if not out:
        raise ValueError("BOM 无有效明细行")

    meta["product_code"] = product_code
    meta["product_name"] = product_name
    meta["line_count"] = len(out)
    return out, meta


def parse_quote_bom_file(path: Union[str, Path]) -> tuple[list[BomInputLine], dict]:
    p = Path(path)
    return parse_quote_bom_bytes(p.read_bytes(), p.name)
