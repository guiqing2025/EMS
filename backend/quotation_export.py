"""导出鼎雄风格加工报价单 Excel"""
from __future__ import annotations

import io
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from sqlalchemy.orm import Session

from models import QuoteBomLine, QuoteCostLine, QuoteOrder


def _thin():
    s = Side(style="thin", color="000000")
    return Border(left=s, right=s, top=s, bottom=s)


def _cost_map(lines: list[QuoteCostLine]) -> dict[str, QuoteCostLine]:
    return {x.item_key: x for x in lines}


def export_quote_xlsx(db: Session, quote: QuoteOrder) -> bytes:
    costs = (
        db.query(QuoteCostLine)
        .filter(QuoteCostLine.quote_id == quote.id)
        .order_by(QuoteCostLine.sort_order)
        .all()
    )
    bom = (
        db.query(QuoteBomLine)
        .filter(QuoteBomLine.quote_id == quote.id)
        .order_by(QuoteBomLine.sort_order)
        .all()
    )
    cm = _cost_map(costs)

    wb = Workbook()
    ws = wb.active
    ws.title = "鼎雄报价模版"

    thin = _thin()
    title_font = Font(name="微软雅黑", size=16, bold=True)
    head_font = Font(name="微软雅黑", size=11, bold=True)
    normal = Font(name="微软雅黑", size=10)

    ws.merge_cells("A1:H1")
    ws["A1"] = "深圳鼎雄电子科技有限公司"
    ws["A1"].font = title_font
    ws["A1"].alignment = Alignment(horizontal="center")

    ws["A2"] = "客户名称"
    ws["B2"] = quote.customer_name or ""
    ws["E2"] = "产品名称"
    ws["F2"] = quote.product_code or quote.product_name or ""

    ws.merge_cells("A3:H3")
    ws["A3"] = "PCBA加工报价单"
    ws["A3"].font = head_font
    ws["A3"].alignment = Alignment(horizontal="center")

    headers = ["工艺项目", "加工内容", "数量", "折算点数", "标准单价", "单位", "金额", "计价说明"]
    for i, h in enumerate(headers, 1):
        cell = ws.cell(4, i, h)
        cell.font = head_font
        cell.border = thin

    def put_row(r: int, section: str, name: str, key: str, note: str = ""):
        row = cm.get(key)
        ws.cell(r, 1, section).border = thin
        ws.cell(r, 2, name).border = thin
        ws.cell(r, 3, row.qty if row else 0).border = thin
        ws.cell(r, 4, row.points if row else 0).border = thin
        ws.cell(r, 5, row.unit_price if row else 0).border = thin
        ws.cell(r, 6, row.unit if row else "").border = thin
        ws.cell(r, 7, row.amount if row else 0).border = thin
        ws.cell(r, 8, (row.note if row and row.note else note)).border = thin
        for c in range(1, 9):
            ws.cell(r, c).font = normal

    put_row(5, "无铅SMT", "CHIP件一类", "chip1", "0603/1206，1颗=1点")
    put_row(6, "", "CHIP件二类", "chip2", "其余贴片，1颗=1.5点")
    put_row(7, "", "异形器件", "odd")
    put_row(8, "", "IC件一（≤8pin）", "ic1")
    put_row(9, "", "IC件二（＞8pin）", "ic2", "P列金额合计")
    put_row(10, "", "大焊盘器件", "bigpad", "TO-263")
    put_row(11, "", "塘锡道上锡", "tangxi")
    put_row(12, "", "SMT小计", "smt_subtotal")

    put_row(13, "无铅插件", "器件二", "dip2", "Q列脚数当点")
    put_row(14, "", "插件小计", "dip_subtotal")

    put_row(15, "后加工", "后加工小计", "post_subtotal")

    ws.cell(16, 1, "产品加工含13%增值税单价(RMB)")
    ws.merge_cells("A16:F16")
    ws.cell(16, 7, cm["unit_price"].amount if cm.get("unit_price") else quote.unit_price)
    for c in range(1, 9):
        ws.cell(16, c).border = thin
        ws.cell(16, c).font = head_font

    put_row(17, "专用治具", "SMT 钢网", "stencil")
    put_row(18, "", "波峰焊治具", "wave_fixture")
    put_row(19, "", "专用治具合计", "tooling_subtotal")

    ws.cell(20, 1, "试产/工程费(RMB)")
    ws.cell(20, 5, quote.engineering_fee or 0)
    ws.cell(21, 1, "批量数量")
    ws.cell(21, 5, quote.batch_qty or 0)
    ws.cell(22, 1, "合计总价(RMB)")
    ws.cell(22, 5, quote.grand_total or 0)
    for r in (20, 21, 22):
        for c in range(1, 9):
            ws.cell(r, c).border = thin

    ws.cell(24, 1, "报价单号")
    ws.cell(24, 2, quote.quote_no)
    ws.cell(25, 1, "其它说明")
    ws.cell(26, 1, "1，批次加工费用=产品加工单价×批量数+专用治具费用+试产/工程费；")
    ws.cell(27, 1, "2，本单按鼎雄报价习惯算法自动核算（CHIP一类0603/1206；二类1.5点；IC取P列；DIP取Q列）。")
    ws.cell(28, 1, f"源文件：{quote.source_filename or '-'}；制单：{quote.created_by or '-'}")

    ws.column_dimensions["A"].width = 18
    ws.column_dimensions["B"].width = 18
    ws.column_dimensions["H"].width = 36

    # BOM 明细 sheet
    ws2 = wb.create_sheet("BOM")
    bom_headers = [
        "序号",
        "元件料号",
        "元件品名",
        "元件规格",
        "单位用量",
        "单位",
        "位号",
        "封装",
        "归类",
        "IC金额P",
        "DIP点数Q",
        "折点",
        "金额",
    ]
    for i, h in enumerate(bom_headers, 1):
        ws2.cell(1, i, h).font = head_font
    for i, line in enumerate(bom, 2):
        vals = [
            line.seq,
            line.material_code,
            line.material_name,
            line.spec,
            line.qty_per,
            line.unit,
            line.position,
            line.package,
            line.bucket,
            line.ic_amount,
            line.dip_points,
            line.calc_points,
            line.calc_amount,
        ]
        for c, v in enumerate(vals, 1):
            ws2.cell(i, c, v)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
