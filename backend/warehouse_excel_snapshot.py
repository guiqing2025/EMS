"""进销表完整快照：按单元格原样归档每个工作表。"""
from __future__ import annotations

import json
import logging
from datetime import datetime
from pathlib import Path
from typing import Any, Optional

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from models import WarehouseSheetSnapshot, WarehouseWorkbookSnapshot

logger = logging.getLogger(__name__)


def _cell_to_json(val: Any) -> Any:
    if val is None:
        return None
    if isinstance(val, datetime):
        return val.strftime("%Y-%m-%d %H:%M:%S") if val.time().second or val.time().minute or val.time().hour else val.strftime("%Y-%m-%d")
    if isinstance(val, bool):
        return val
    if isinstance(val, (int, float)):
        return val
    text = str(val)
    return text


def _sheet_to_grid(ws) -> tuple[list[list[Any]], int, int]:
    max_row = ws.max_row or 0
    max_col = ws.max_column or 0
    if max_row <= 0 or max_col <= 0:
        return [], 0, 0
    grid: list[list[Any]] = []
    for r in range(1, max_row + 1):
        row_vals = [_cell_to_json(ws.cell(r, c).value) for c in range(1, max_col + 1)]
        grid.append(row_vals)
    return grid, max_row, max_col


def archive_workbook_snapshot(
    db: Session,
    file_path: Path,
    customer_id: str,
    customer_name: str,
) -> dict[str, int]:
    """将 workbook 每个 sheet 的全部单元格写入数据库（显示值，结构与 Excel 一致）。"""
    path = Path(file_path)
    wb = load_workbook(path, read_only=False, data_only=True)
    try:
        file_mtime = path.stat().st_mtime
        db.query(WarehouseSheetSnapshot).filter(
            WarehouseSheetSnapshot.customer_id == customer_id,
            WarehouseSheetSnapshot.source_file == str(path),
        ).delete(synchronize_session=False)
        db.query(WarehouseWorkbookSnapshot).filter(
            WarehouseWorkbookSnapshot.customer_id == customer_id,
            WarehouseWorkbookSnapshot.source_file == str(path),
        ).delete(synchronize_session=False)

        snapshot = WarehouseWorkbookSnapshot(
            customer_id=customer_id,
            customer_name=customer_name,
            source_file=str(path),
            file_name=path.name,
            file_mtime=file_mtime,
            sheet_names=json.dumps(wb.sheetnames, ensure_ascii=False),
            sheet_count=len(wb.sheetnames),
        )
        db.add(snapshot)
        db.flush()

        total_cells = 0
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            grid, row_count, col_count = _sheet_to_grid(ws)
            total_cells += row_count * col_count
            db.add(
                WarehouseSheetSnapshot(
                    snapshot_id=snapshot.id,
                    customer_id=customer_id,
                    customer_name=customer_name,
                    source_file=str(path),
                    sheet_name=sheet_name,
                    row_count=row_count,
                    col_count=col_count,
                    data_json=json.dumps(grid, ensure_ascii=False),
                )
            )
        return {
            "snapshot_id": snapshot.id,
            "sheet_count": len(wb.sheetnames),
            "total_cells": total_cells,
        }
    finally:
        wb.close()


def archive_workbook_snapshot_from_path(
    db: Session,
    file_path: Path,
    customer_id: str,
    customer_name: str,
) -> Optional[dict[str, int]]:
    try:
        return archive_workbook_snapshot(db, file_path, customer_id, customer_name)
    except Exception:
        logger.exception("进销表快照失败 %s", file_path)
        return None
