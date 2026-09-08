"""共享盘 Excel 物料同步（读取客户进销表/库存表）"""
from __future__ import annotations

import logging
import re
from datetime import datetime
from pathlib import Path
from typing import Optional

from openpyxl import Workbook, load_workbook
from sqlalchemy.exc import OperationalError
from sqlalchemy.orm import Session

from config import get_customer, get_warehouse_share_path, load_config
from models import ExcelImportLog, WarehouseMaterial

logger = logging.getLogger(__name__)

TEMPLATE_HEADERS = ["物料编码", "物料名称", "规格型号", "单位", "期初库存", "备注"]
SHEET_NAME = "物料明细"
INVENTORY_SHEET = "库存表"

SHARE_DIR_CANDIDATES = [
    lambda: Path(get_warehouse_share_path()),
    lambda: Path.home() / "Desktop/共享-测试软件资料/D-仓库表格/客户进销表",
]


def _path_is_dir_quick(path: Path, timeout_sec: float = 1.2) -> bool:
    """共享盘可能未挂载导致 is_dir 挂起，限时探测。

    同时尝试 listdir：macOS TCC 下 is_dir 可能成功但无权遍历，
    仅测 is_dir 会误报「已连接」。
    """
    import threading

    box: dict[str, bool] = {"ok": False}

    def _check() -> None:
        try:
            if not path.is_dir():
                box["ok"] = False
                return
            next(path.iterdir(), None)
            box["ok"] = True
        except Exception:
            box["ok"] = False

    t = threading.Thread(target=_check, daemon=True)
    t.start()
    t.join(timeout_sec)
    return bool(box["ok"]) if not t.is_alive() else False


def resolve_share_dir() -> Path:
    for factory in SHARE_DIR_CANDIDATES:
        path = factory()
        if _path_is_dir_quick(path):
            return path
    return Path(get_warehouse_share_path())


SKIP_FOLDERS = {"超领单", "仓库过往资料"}

# 仅同步这 8 个客户文件夹（与共享盘目录名一致）
ALLOWED_CUSTOMER_FOLDERS = frozenset(
    {
        "恩玖进销表",
        "菲利斯进销表",
        "华夏恒泰进销表",
        "能系科技",
        "亿兰科进销表",
        "亿维艾",
        "永联发料单",
        "源信进销表",
    }
)


def _is_customer_folder(name: str) -> bool:
    if name in SKIP_FOLDERS or name.startswith("."):
        return False
    return name in ALLOWED_CUSTOMER_FOLDERS


def _match_customer(folder_name: str) -> tuple[str, str]:
    name = folder_name.strip()
    aliases = {
        "菲利斯进销表": ("feilisi", "菲利斯"),
        "恩玖进销表": ("enjiu", "恩玖·鼎雄"),
        "华夏恒泰进销表": ("wh_huaxia", "华夏恒泰"),
        "永联发料单": ("yonglian", "永联"),
        "源信进销表": ("wh_yuanxin", "源信"),
        "亿兰科进销表": ("wh_yilanke", "亿兰科"),
        "亿维艾": ("wh_yiweiai", "亿维艾"),
        "能系科技": ("wh_nengxi", "能系科技"),
    }
    if name in aliases:
        cid, cname = aliases[name]
        customer = get_customer(cid)
        if customer:
            return cid, customer.get("name") or cname
        return cid, cname
    short = re.sub(r"(进销表|发料单)$", "", name)
    return f"wh_{short}", short or name


def _find_latest_workbook(folder: Path) -> Optional[Path]:
    files = [
        p
        for p in folder.iterdir()
        if p.is_file()
        and p.suffix.lower() in (".xlsx", ".xlsm")
        and not p.name.startswith("~$")
        and not p.name.startswith(".")
    ]
    if not files:
        return None
    files.sort(key=lambda p: p.stat().st_mtime, reverse=True)
    return files[0]


def _header_map(header_row: tuple) -> dict[str, int]:
    mapping: dict[str, int] = {}
    for idx, cell in enumerate(header_row):
        if cell is None:
            continue
        key = str(cell).strip()
        if key:
            mapping[key] = idx
    return mapping


def _cell(row: tuple, mapping: dict[str, int], *names: str):
    for name in names:
        idx = mapping.get(name)
        if idx is not None and idx < len(row):
            val = row[idx]
            if val not in (None, ""):
                return val
    return None


def _float_cell(row: tuple, mapping: dict[str, int], *names: str) -> Optional[float]:
    val = _cell(row, mapping, *names)
    if val in (None, ""):
        return None
    try:
        return float(val)
    except (TypeError, ValueError):
        return None


def _parse_inventory_rows(ws) -> list[dict]:
    """只取库存表：物料编号 + 实际库存（数）。"""
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    header_map: dict[str, int] = {}
    for idx, row in enumerate(rows[:20]):
        if not row:
            continue
        cells = [str(c).strip() if c is not None else "" for c in row]
        if "物料编号" in cells or "物料编码" in cells:
            header_idx = idx
            header_map = _header_map(row)
            break
    if header_idx is None:
        raise ValueError("库存表未找到「物料编号/物料编码」表头")
    if "实际库存数" not in header_map and "实际库存" not in header_map:
        raise ValueError("库存表缺少「实际库存」或「实际库存数」列")

    items: list[dict] = []
    for row in rows[header_idx + 1 :]:
        if not row:
            continue
        code = _cell(row, header_map, "物料编号", "物料编码")
        code = str(code or "").strip()
        if not code or code in ("物料编号", "物料编码"):
            continue
        actual = _float_cell(row, header_map, "实际库存数", "实际库存")
        if actual is None:
            actual = 0.0
        items.append(
            {
                "material_code": code,
                "material_name": str(_cell(row, header_map, "物料名称") or "").strip(),
                "spec": str(_cell(row, header_map, "规格型号") or "").strip(),
                "unit": "PCS",
                "opening_qty": actual,
                "excel_count_qty": actual,
                "excel_in_qty": None,
                "excel_demand_qty": _float_cell(row, header_map, "需求数"),
                "remark": "",
            }
        )
    return items


def _parse_template_rows(ws) -> list[dict]:
    items: list[dict] = []
    for idx, row in enumerate(ws.iter_rows(values_only=True)):
        if idx == 0:
            continue
        code = str(row[0] or "").strip() if len(row) > 0 and row[0] is not None else ""
        if not code or code == "物料编码":
            continue
        opening = 0
        if len(row) > 4 and row[4] not in (None, ""):
            try:
                opening = float(row[4])
            except (TypeError, ValueError):
                opening = 0
        items.append(
            {
                "material_code": code,
                "material_name": str(row[1] or "").strip() if len(row) > 1 else "",
                "spec": str(row[2] or "").strip() if len(row) > 2 else "",
                "unit": (str(row[3] or "PCS").strip() if len(row) > 3 else "PCS") or "PCS",
                "opening_qty": opening,
                "remark": str(row[5] or "").strip() if len(row) > 5 else "",
            }
        )
    return items


def upsert_material(
    db: Session,
    customer_id: str,
    customer_name: str,
    item: dict,
    *,
    sync_stock: bool = False,
) -> tuple[bool, WarehouseMaterial]:
    row = (
        db.query(WarehouseMaterial)
        .filter(
            WarehouseMaterial.customer_id == customer_id,
            WarehouseMaterial.material_code == item["material_code"],
        )
        .first()
    )
    created = row is None
    if not row:
        row = WarehouseMaterial(
            customer_id=customer_id,
            customer_name=customer_name,
            material_code=item["material_code"],
            qty=0,
            locked_qty=0,
        )
        db.add(row)
    row.material_name = item.get("material_name") or row.material_name
    row.spec = item.get("spec") or row.spec
    row.unit = item.get("unit") or row.unit or "PCS"
    row.remark = item.get("remark") or row.remark
    row.customer_name = customer_name
    if sync_stock:
        if "excel_count_qty" in item:
            row.excel_count_qty = item.get("excel_count_qty")
        if "excel_in_qty" in item:
            row.excel_in_qty = item.get("excel_in_qty")
        if "excel_demand_qty" in item:
            row.excel_demand_qty = item.get("excel_demand_qty")
        row.excel_synced_at = datetime.utcnow()
    opening = float(item.get("opening_qty") or 0)
    if created:
        # 新建料号时可用共享盘期初；已有料号的 qty 一律以系统录入为准，禁止同步覆盖
        row.qty = opening
        if opening > 0:
            db.flush()
            from warehouse_service import record_inbound

            record_inbound(
                db,
                row,
                opening,
                "excel_sync",
                operator="系统同步",
                remark=f"共享盘首次同步 · {customer_name}",
                ref_no=f"SYNC-{customer_id}-{item['material_code'][:20]}",
            )
    # sync_stock 仅更新 excel_* 参考字段，不再改 row.qty（避免冲掉系统进账）
    row.updated_at = datetime.utcnow()
    return created, row


def _consolidate_substitute_inventory(items: list[dict]) -> list[dict]:
    """将客户替代料号库存合并到我司主料号，避免同物两行账。"""
    from substitution_service import get_stock_primary_code, normalize_code

    buckets: dict[str, dict] = {}
    alias_notes: dict[str, list[str]] = {}
    for item in items:
        code = normalize_code(item.get("material_code"))
        if not code:
            continue
        primary = get_stock_primary_code(code)
        target = primary or code
        qty = float(item.get("opening_qty") or 0)
        count_qty = float(item.get("excel_count_qty") or 0)
        in_qty = float(item.get("excel_in_qty") or 0)
        if target not in buckets:
            buckets[target] = {
                **item,
                "material_code": target,
                "opening_qty": 0.0,
                "excel_count_qty": 0.0,
                "excel_in_qty": 0.0,
            }
        else:
            bucket = buckets[target]
            bucket["material_name"] = bucket.get("material_name") or item.get("material_name")
            bucket["spec"] = bucket.get("spec") or item.get("spec")
        buckets[target]["opening_qty"] = float(buckets[target].get("opening_qty") or 0) + qty
        buckets[target]["excel_count_qty"] = float(buckets[target].get("excel_count_qty") or 0) + count_qty
        buckets[target]["excel_in_qty"] = float(buckets[target].get("excel_in_qty") or 0) + in_qty
        if code != target and qty:
            alias_notes.setdefault(target, []).append(f"{code}:{qty:g}")
    for target, notes in alias_notes.items():
        if target in buckets:
            extra = "；".join(notes)
            remark = (buckets[target].get("remark") or "").strip()
            buckets[target]["remark"] = f"{remark}；替代料并入 {extra}".strip("；")
    result = list(buckets.values())
    seen = {normalize_code(item["material_code"]) for item in result}
    for item in items:
        code = normalize_code(item.get("material_code"))
        primary = get_stock_primary_code(code)
        if code and primary and code != primary and code not in seen:
            result.append(
                {
                    "material_code": code,
                    "material_name": item.get("material_name") or "",
                    "spec": item.get("spec") or "",
                    "unit": item.get("unit") or "PCS",
                    "opening_qty": 0.0,
                    "remark": f"库存已并入主料 {primary}",
                }
            )
    return result


def import_workbook(db: Session, file_path: Path, customer_id: str, customer_name: str) -> ExcelImportLog:
    """只同步库存表「实际库存」到 warehouse_materials；不跑进出账全量扫描。"""
    imported = updated = 0
    try:
        wb = load_workbook(file_path, read_only=True, data_only=True)
        try:
            if INVENTORY_SHEET not in wb.sheetnames:
                raise ValueError(f"缺少「{INVENTORY_SHEET}」工作表")
            ws = wb[INVENTORY_SHEET]
            items = _parse_inventory_rows(ws)
            items = _consolidate_substitute_inventory(items)
            for item in items:
                created, _ = upsert_material(
                    db,
                    customer_id,
                    customer_name,
                    item,
                    sync_stock=True,
                )
                if created:
                    imported += 1
                else:
                    updated += 1
        finally:
            wb.close()

        msg = f"库存表实际库存：读取 {len(items)} 行，新增 {imported}，更新 {updated}"
        log = ExcelImportLog(
            source_file=str(file_path),
            customer_name=customer_name,
            rows_imported=imported,
            rows_updated=updated,
            status="success",
            message=msg,
        )
    except Exception as exc:
        logger.exception("同步失败 %s", file_path)
        try:
            db.rollback()
        except Exception:
            pass
        msg = str(exc)
        if "database is locked" in msg.lower() or "locked" in msg.lower():
            msg = "数据库正忙（可能 ICT/AOI 同步中），请稍后再点「从共享盘导入」"
        elif "UNIQUE constraint failed: warehouse_movements.dedupe_key" in msg:
            msg = "进出账重复键冲突（换月文件与旧流水撞车），请重试；若仍失败请联系管理员"
        log = ExcelImportLog(
            source_file=str(file_path),
            customer_name=customer_name,
            status="failed",
            message=msg,
        )
    db.add(log)
    return log


def _retry_db_locked(db: Session, fn, *, attempts: int = 8, base_sleep: float = 0.5):
    """SQLite 被 ICT/AOI 占锁时：用 SAVEPOINT 重试，避免 rollback 冲掉已解析的物料行。"""
    import time
    from typing import Optional

    last: Optional[Exception] = None
    for i in range(attempts):
        try:
            with db.begin_nested():
                return fn()
        except OperationalError as exc:
            last = exc
            if "locked" not in str(exc).lower():
                raise
            time.sleep(base_sleep * (i + 1))
    assert last is not None
    raise last


def sync_all_from_share(db: Session) -> tuple[list[ExcelImportLog], str]:
    share_dir = resolve_share_dir()
    if not share_dir.is_dir():
        raise FileNotFoundError(
            f"共享盘目录不可访问: {share_dir}。"
            f"请确认 Mac 已挂载 192.168.2.11（当前配置: {get_warehouse_share_path()}）"
        )

    logs: list[ExcelImportLog] = []

    # 客户子文件夹：菲利斯进销表、恩玖进销表 …
    try:
        children = sorted(share_dir.iterdir())
    except OSError as exc:
        raise FileNotFoundError(
            f"共享盘目录无权读取: {share_dir}（{exc}）。"
            f"请用本机终端重启后端，并确认已授予「桌面文件夹」或「完全磁盘访问」权限"
        ) from exc

    for child in children:
        if not child.is_dir() or child.name.startswith(".") or not _is_customer_folder(child.name):
            continue
        book = _find_latest_workbook(child)
        if not book:
            continue
        customer_id, customer_name = _match_customer(child.name)
        log = import_workbook(db, book, customer_id, customer_name)
        logs.append(log)
        # 每客户提交一次，缩短锁持有时间，避免拖死整次导入
        try:
            db.commit()
        except Exception:
            logger.exception("提交客户同步失败 %s", child.name)
            db.rollback()
            if log.status == "success":
                log.status = "failed"
                log.message = "提交失败（数据库忙），请重试"
                db.add(log)
                try:
                    db.commit()
                except Exception:
                    db.rollback()

    if not logs:
        raise FileNotFoundError(f"未在共享盘找到可同步的客户进销表: {share_dir}")

    ok = sum(1 for log in logs if log.status == "success")
    summary = f"已从 {share_dir} 同步 {ok}/{len(logs)} 个客户文件"
    failed = [log for log in logs if log.status != "success"]
    if failed:
        names = "、".join((log.customer_name or "?") for log in failed[:5])
        more = f"等{len(failed)}个" if len(failed) > 5 else ""
        tip = (failed[0].message or "")[:80]
        summary += f"；失败：{names}{more}"
        if tip:
            summary += f"（例：{tip}）"
    return logs, summary


def build_template_workbook() -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(TEMPLATE_HEADERS)
    ws.append(["示例料号", "示例品名", "规格", "PCS", 100, "备注可选"])
    return wb


def export_inventory_workbook(db: Session, customer_id: Optional[str] = None) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = SHEET_NAME
    ws.append(["客户", *TEMPLATE_HEADERS, "来料数", "需求数", "当前库存", "锁定库存"])
    q = db.query(WarehouseMaterial).order_by(
        WarehouseMaterial.customer_name, WarehouseMaterial.material_code
    )
    if customer_id:
        q = q.filter(WarehouseMaterial.customer_id == customer_id)
    for row in q.all():
        ws.append([
            row.customer_name,
            row.material_code,
            row.material_name or "",
            row.spec or "",
            row.unit or "PCS",
            "",
            row.remark or "",
            row.excel_in_qty,
            row.excel_demand_qty,
            row.qty,
            row.locked_qty,
        ])
    return wb
